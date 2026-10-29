import hashlib
import hmac
import json
import os
import re
import secrets
import smtplib
import ssl
import traceback
import warnings
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from email.message import EmailMessage

from fastapi import BackgroundTasks, Depends, FastAPI, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, inspect, or_, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload
from starlette.middleware.sessions import SessionMiddleware

import models
from database import BASE_DIR, SessionLocal, engine
from excel_com_fm_08 import (
    generar_excel_com_fm_08,
    _escribir_bloque_po_en_hoja,
    _preparar_hoja_cuadro,
    generar_excel_consolidado_usuario,
    generar_excel_maestro_productos,
    COLORES_FILAS_PO
)

warnings.filterwarnings("ignore", message="Dialect sqlite.*Decimal")

CLAVE_MAESTRA = os.getenv("CLAVE_MAESTRA", "FLY01")
CORREO_OFICIAL = "notificaciones@flylogistic.aero"
CLAVE_CORREO_PRUEBA = "Colombia2024*"
NOMBRE_ADMIN = "Juan Esteban Martínez"

ROLES_SISTEMA = ["Administrador", "Gerencia", "Jefe de area", "Asesor"]
ROLES_SUPERIORES = ["Administrador", "Gerencia", "Jefe de area"]
ROLES_ACCESO_USUARIOS = ["Administrador", "Gerencia", "Jefe de area"]

ESTADOS = ["En proceso", "Entregado"]
ESTADOS_ADUANA = ["No aplica / Pendiente", "En revisión aduanera", "Pago de aranceles", "Liberado / Nacionalizado"]
ESTADOS_PRODUCTO = ["PENDIENTE", "DESPACHADO", "EN FLY", "LLEGADO FLY", "ENTREGADO"]
METODOS_PAGO_OPCIONES = ["Visa", "American", "PayPal", "Transferencia"]
UNIDADES_MEDIDA = ["UND", "EA", "GAL", "LT", "KIT", "LB", "ROLLO", "CAJA", "MT", "PC", "SET"]

PESOS_PROBABILIDAD = {
    "En proceso": 25,
    "Entregado": 100,
}


def cargar_secret_key() -> str:
    env = os.getenv("SECRET_KEY")
    if env:
        return env
    ruta = os.path.join(BASE_DIR, ".secret_key")
    if os.path.exists(ruta):
        with open(ruta) as f:
            return f.read().strip()
    clave = secrets.token_hex(32)
    with open(ruta, "w") as f:
        f.write(clave)
    return clave


GOOGLE_DRIVE_DISPONIBLE = False
try:
    from google.auth.transport.requests import Request as GoogleRequest
    from google.oauth2.credentials import Credentials
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload
    GOOGLE_DRIVE_DISPONIBLE = True
except ImportError:
    pass

SCOPES = ["https://www.googleapis.com/auth/drive.file"]
TOKEN_PATH = os.path.join(BASE_DIR, "token.json")
FOLDER_MIME = "application/vnd.google-apps.folder"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_drive_service = None


def obtener_drive():
    global _drive_service
    if not GOOGLE_DRIVE_DISPONIBLE:
        return None
    if _drive_service:
        return _drive_service
    if not os.path.exists(TOKEN_PATH):
        return None
    creds = Credentials.from_authorized_user_file(TOKEN_PATH, SCOPES)
    if not creds.valid:
        if creds.expired and creds.refresh_token:
            creds.refresh(GoogleRequest())
            with open(TOKEN_PATH, "w") as f:
                f.write(creds.to_json())
        else:
            return None
    _drive_service = build("drive", "v3", credentials=creds, cache_discovery=False)
    return _drive_service


def _q(texto: str) -> str:
    return texto.replace("\\", "\\\\").replace("'", "\\'")


def carpeta_drive(service, nombre, padre_id=None):
    q = f"name='{_q(nombre)}' and mimeType='{FOLDER_MIME}' and trashed=false"
    if padre_id:
        q += f" and '{padre_id}' in parents"
    res = service.files().list(q=q, fields="files(id, webViewLink)", pageSize=1).execute()
    if res["files"]:
        return res["files"][0]["id"], res["files"][0].get("webViewLink", "")
    meta = {"name": nombre, "mimeType": FOLDER_MIME}
    if padre_id:
        meta["parents"] = [padre_id]
    creada = service.files().create(body=meta, fields="id, webViewLink").execute()
    return creada["id"], creada.get("webViewLink", "")


def subir_archivo_a_carpeta_drive(service, ruta_local, nombre_archivo, id_carpeta, mime_type=None):
    media = MediaFileUpload(ruta_local, mimetype=mime_type, resumable=False)
    existente = service.files().list(
        q=f"name='{_q(nombre_archivo)}' and '{id_carpeta}' in parents and trashed=false",
        fields="files(id)", pageSize=1
    ).execute()["files"]
    if existente:
        return service.files().update(
            fileId=existente[0]["id"], media_body=media, fields="id, webViewLink"
        ).execute()
    return service.files().create(
        body={"name": nombre_archivo, "parents": [id_carpeta]},
        media_body=media, fields="id, webViewLink"
    ).execute()


def sincronizar_po_en_drive(service, p: models.Pedido, ruta_excel_local: str, nombre_excel: str, ruta_doc_local: str | None = None):
    raiz = "Pedidos Nacionales" if p.tipo == "Nacional" else "Pedidos Internacionales"
    id_raiz, _ = carpeta_drive(service, raiz)
    id_asesor, _ = carpeta_drive(service, p.asesor.nombre, id_raiz)
    id_po_folder, link_po_folder = carpeta_drive(service, f"PO_{p.po}", id_asesor)

    res_excel = subir_archivo_a_carpeta_drive(service, ruta_excel_local, nombre_excel, id_po_folder, XLSX_MIME)
    link_excel = res_excel.get("webViewLink", "")

    link_doc = p.documentacion_drive_link or ""
    if ruta_doc_local and os.path.exists(ruta_doc_local):
        nombre_doc = os.path.basename(ruta_doc_local)
        res_doc = subir_archivo_a_carpeta_drive(service, ruta_doc_local, nombre_doc, id_po_folder)
        link_doc = res_doc.get("webViewLink", "")

    return link_excel, link_po_folder, link_doc


def nombre_seguro(texto: str) -> str:
    s = re.sub(r"[^\w\-. ]", "_", (texto or "").strip(), flags=re.UNICODE).strip(" .")
    return s or "sin_nombre"


def carpeta_po_local(p: models.Pedido) -> str:
    raiz_local = "Pedidos Nacionales" if p.tipo == "Nacional" else "Pedidos Internacionales"
    nombre_asesor = p.asesor.nombre if p.asesor else "General"
    carpeta = os.path.join(
        BASE_DIR,
        raiz_local,
        nombre_seguro(nombre_asesor),
        f"PO_{nombre_seguro(p.po)}"
    )
    os.makedirs(carpeta, exist_ok=True)
    return carpeta


def ruta_excel(p: models.Pedido):
    carpeta = carpeta_po_local(p)
    nombre = f"Pedido_{nombre_seguro(p.po)}.xlsx"
    return os.path.join(carpeta, nombre), nombre


async def guardar_archivo_documentacion(p: models.Pedido, archivo_subido: UploadFile | None) -> str | None:
    if not archivo_subido or not archivo_subido.filename or not archivo_subido.filename.strip():
        if p.documentacion_archivo:
            posible = os.path.join(carpeta_po_local(p), p.documentacion_archivo)
            if os.path.exists(posible):
                return posible
        return None
    nombre_limpio = nombre_seguro(os.path.basename(archivo_subido.filename))
    carpeta = carpeta_po_local(p)
    ruta_destino = os.path.join(carpeta, nombre_limpio)
    contenido = await archivo_subido.read()
    if contenido:
        with open(ruta_destino, "wb") as f_out:
            f_out.write(contenido)
        p.documentacion_archivo = nombre_limpio
        if not p.documentacion:
            p.documentacion = nombre_limpio
        return ruta_destino
    return None


def generar_salidas(db: Session, p: models.Pedido, ruta_doc_local: str | None = None) -> str:
    try:
        ruta, nombre = ruta_excel(p)
        generar_excel_com_fm_08(p, ruta)
    except Exception as e:
        return f"Orden guardada, error al generar Excel: {e}"
    try:
        service = obtener_drive()
        if service:
            if not ruta_doc_local and p.documentacion_archivo:
                posible_doc = os.path.join(carpeta_po_local(p), p.documentacion_archivo)
                if os.path.exists(posible_doc):
                    ruta_doc_local = posible_doc
            link_excel, link_folder, link_doc = sincronizar_po_en_drive(service, p, ruta, nombre, ruta_doc_local)
            p.drive_link = link_excel or link_folder
            p.drive_folder_link = link_folder
            if link_doc:
                p.documentacion_drive_link = link_doc
            db.commit()
            return "Expediente y Excel sincronizados en Google Drive"
        return "Expediente y Excel generados correctamente"
    except Exception as e:
        return f"Excel generado localmente ({e})"


def sugerir_siguiente_po(db: Session) -> str:
    todas_pos = [r[0] for r in db.query(models.Pedido.po).all() if r[0]]
    max_num = 0
    prefijo_detectado = "PO-"
    ancho_num = 4
    for po_str in todas_pos:
        m = re.search(r"^(.*?)(\d+)$", po_str.strip())
        if m:
            num_val = int(m.group(2))
            if num_val > max_num:
                max_num = num_val
                prefijo_detectado = m.group(1)
                ancho_num = len(m.group(2))
    if max_num == 0:
        return "PO-0001"
    return f"{prefijo_detectado}{str(max_num + 1).zfill(ancho_num)}"


ITER = 200_000


def hash_clave(clave: str) -> str:
    sal = secrets.token_hex(16)
    h = hashlib.pbkdf2_hmac("sha256", clave.encode(), bytes.fromhex(sal), ITER).hex()
    return f"pbkdf2${ITER}${sal}${h}"


def verificar_clave(clave: str, guardado) -> bool:
    try:
        _, iters, sal, h = (guardado or "").split("$")
        calc = hashlib.pbkdf2_hmac("sha256", clave.encode(), bytes.fromhex(sal), int(iters)).hex()
        return hmac.compare_digest(calc, h)
    except ValueError:
        hashlib.pbkdf2_hmac("sha256", clave.encode(), b"x" * 16, ITER)
        return False


def buscar_autorizador_por_clave(db: Session, clave_ingresada: str, jefe_id_seleccionado: int | None = None):
    if not clave_ingresada:
        return False, None
    if hmac.compare_digest(clave_ingresada.strip().encode(), CLAVE_MAESTRA.encode()):
        jefe_sel = db.get(models.Asesor, jefe_id_seleccionado) if jefe_id_seleccionado else None
        return True, jefe_sel
    superiores = db.query(models.Asesor).filter(models.Asesor.rol.in_(ROLES_SUPERIORES)).all()
    if jefe_id_seleccionado:
        jefe_sel = db.get(models.Asesor, jefe_id_seleccionado)
        if jefe_sel and jefe_sel.password_hash and verificar_clave(clave_ingresada, jefe_sel.password_hash):
            return True, jefe_sel
    for sup in superiores:
        if sup.password_hash and verificar_clave(clave_ingresada, sup.password_hash):
            return True, sup
    return False, None


def migrar_bd():
    insp = inspect(engine)
    with engine.begin() as con:
        for tabla in models.Base.metadata.sorted_tables:
            if not insp.has_table(tabla.name):
                continue
            existentes = {c["name"] for c in insp.get_columns(tabla.name)}
            for col in tabla.columns:
                if col.name not in existentes:
                    tipo = col.type.compile(engine.dialect)
                    con.execute(text(f'ALTER TABLE {tabla.name} ADD COLUMN "{col.name}" {tipo}'))

        if insp.has_table("asesores"):
            con.execute(text("UPDATE asesores SET rol='Asesor' WHERE rol IS NULL OR rol='' OR rol='Supervisión'"))
            con.execute(text("UPDATE asesores SET email=:correo WHERE email IS NULL OR email=''"),
                        {"correo": CORREO_OFICIAL})
            filas = con.execute(text(
                "SELECT id, nombre FROM asesores WHERE nombre_norm IS NULL OR nombre_norm = ''")).fetchall()
            for i, nombre in filas:
                con.execute(text("UPDATE asesores SET nombre_norm=:n WHERE id=:i"),
                            {"n": models.normalizar(nombre), "i": i})
            con.execute(text("CREATE UNIQUE INDEX IF NOT EXISTS ix_asesores_username ON asesores(username)"))

        if insp.has_table("pedidos"):
            con.execute(text("UPDATE pedidos SET valor_nacionalizacion = 0 WHERE valor_nacionalizacion IS NULL"))
            con.execute(text("UPDATE pedidos SET estado='En proceso' WHERE estado NOT IN ('En proceso', 'Entregado')"))

        if insp.has_table("pedido_items"):
            con.execute(text("UPDATE pedido_items SET parte_numero_alterno = '' WHERE parte_numero_alterno IS NULL"))
            con.execute(text("UPDATE pedido_items SET unidad_medida = 'UND' WHERE unidad_medida IS NULL OR unidad_medida = ''"))
            con.execute(text("UPDATE pedido_items SET iva_porcentaje = 19 WHERE iva_porcentaje IS NULL"))
            con.execute(text("UPDATE pedido_items SET estado_producto = 'PENDIENTE' WHERE estado_producto IS NULL OR estado_producto = ''"))

        if insp.has_table("pedidos_internacionales"):
            con.execute(text("UPDATE pedidos_internacionales SET llega_click_mail = 'Si' WHERE fecha_click_mail IS NOT NULL AND (llega_click_mail IS NULL OR llega_click_mail = '')"))
            con.execute(text("UPDATE pedidos_internacionales SET llega_click_mail = 'No' WHERE llega_click_mail IS NULL OR llega_click_mail = ''"))


def inicializar_admin_y_correo():
    ruta_principal = os.path.join(BASE_DIR, "correo_config.json")
    if not os.path.exists(ruta_principal):
        datos_smtp = {
            "host": "mail.flylogistic.aero",
            "port": 465,
            "user": CORREO_OFICIAL,
            "password": CLAVE_CORREO_PRUEBA,
            "from": f"Fly Logistic S.A.S. <{CORREO_OFICIAL}>",
            "app_url": "http://127.0.0.1:8000"
        }
        for nombre_cfg in ("correo_config.json", "correo_config.json.json"):
            try:
                with open(os.path.join(BASE_DIR, nombre_cfg), "w", encoding="utf-8") as f:
                    json.dump(datos_smtp, f, indent=2, ensure_ascii=False)
            except Exception:
                pass

    db = SessionLocal()
    try:
        # 1. Administrador General
        n_norm = models.normalizar(NOMBRE_ADMIN)
        admin = db.query(models.Asesor).filter(func.lower(models.Asesor.username) == "admin").first()
        if not admin:
            admin = db.query(models.Asesor).filter_by(nombre_norm=n_norm).first()
        if not admin:
            admin = models.Asesor(
                nombre=NOMBRE_ADMIN,
                nombre_norm=n_norm,
                username="admin",
                email=CORREO_OFICIAL,
                rol="Administrador"
            )
            db.add(admin)
        else:
            admin.nombre = NOMBRE_ADMIN
            admin.nombre_norm = n_norm
            admin.username = "admin"
            admin.rol = "Administrador"
            admin.email = admin.email or CORREO_OFICIAL
        admin.password_hash = hash_clave(CLAVE_CORREO_PRUEBA)
        db.commit()
        db.refresh(admin)

        # 2. Leonardo Garzón (Jefe de área)
        leo_norm = models.normalizar("Leonardo Garzón")
        leonardo = db.query(models.Asesor).filter(
            or_(models.Asesor.username == "lgarzon", models.Asesor.nombre_norm == leo_norm)
        ).first()
        if not leonardo:
            leonardo = models.Asesor(
                nombre="Leonardo Garzón",
                nombre_norm=leo_norm,
                username="lgarzon",
                email="leonardo.garzon@flylogistic.aero",
                rol="Jefe de area",
                jefe_id=admin.id,
                password_hash=hash_clave("Colombia2024*")
            )
            db.add(leonardo)
        else:
            leonardo.nombre = "Leonardo Garzón"
            leonardo.nombre_norm = leo_norm
            leonardo.username = "lgarzon"
            leonardo.rol = "Jefe de area"
            leonardo.jefe_id = admin.id
            leonardo.password_hash = hash_clave("Colombia2024*")
        db.commit()
        db.refresh(leonardo)

        # 3. Miguel Mosquera (Asesor)
        miguel_norm = models.normalizar("Miguel Mosquera")
        miguel = db.query(models.Asesor).filter(
            or_(models.Asesor.username == "mmosquera", models.Asesor.nombre_norm == miguel_norm)
        ).first()
        if not miguel:
            miguel = models.Asesor(
                nombre="Miguel Mosquera",
                nombre_norm=miguel_norm,
                username="mmosquera",
                email="miguel.mosquera@flylogistic.aero",
                rol="Asesor",
                jefe_id=leonardo.id,
                password_hash=hash_clave("Colombia2024*")
            )
            db.add(miguel)
        else:
            miguel.nombre = "Miguel Mosquera"
            miguel.nombre_norm = miguel_norm
            miguel.username = "mmosquera"
            miguel.rol = "Asesor"
            miguel.jefe_id = leonardo.id
            miguel.password_hash = hash_clave("Colombia2024*")
        db.commit()

        # 4. Carlos Mendoza (Asesor)
        carlos_norm = models.normalizar("Carlos Mendoza")
        carlos = db.query(models.Asesor).filter(
            or_(models.Asesor.username == "cmendoza", models.Asesor.nombre_norm == carlos_norm)
        ).first()
        if not carlos:
            carlos = models.Asesor(
                nombre="Carlos Mendoza",
                nombre_norm=carlos_norm,
                username="cmendoza",
                email="carlos.mendoza@flylogistic.aero",
                rol="Asesor",
                jefe_id=leonardo.id,
                password_hash=hash_clave("Colombia2024*")
            )
            db.add(carlos)
        else:
            carlos.nombre = "Carlos Mendoza"
            carlos.nombre_norm = carlos_norm
            carlos.username = "cmendoza"
            carlos.rol = "Asesor"
            carlos.jefe_id = leonardo.id
            carlos.password_hash = hash_clave("Colombia2024*")
        db.commit()

        # Limpiar cuentas antiguas de prueba si existen y pasar sus POs a Leonardo Garzón
        for old_u in db.query(models.Asesor).filter(models.Asesor.username.in_(["qa_tester", "asesor_qa"])).all():
            db.query(models.Pedido).filter(models.Pedido.asesor_id == old_u.id).update(
                {"asesor_id": leonardo.id}, synchronize_session=False
            )
            db.delete(old_u)
        db.commit()

    except Exception as e:
        db.rollback()
        print(f"[AVISO INICIALIZACIÓN]: {e}")
    finally:
        db.close()


models.Base.metadata.create_all(bind=engine)
migrar_bd()
inicializar_admin_y_correo()

app = FastAPI(title="Fly Logistic S.A.S. - OMS")


@app.middleware("http")
async def interceptor_de_errores(request: Request, call_next):
    try:
        return await call_next(request)
    except Exception:
        trace = traceback.format_exc()
        print("\n[ERROR INTERNO DETECTADO]\n", trace)
        return HTMLResponse(
            content=f"""
            <div style="background:#0f172a; color:#38bdf8; padding:32px; font-family:monospace; min-height:100vh;">
                <h2 style="color:#f87171; margin-bottom:12px;">Diagnóstico Técnico del Servidor</h2>
                <p style="color:#e2e8f0; margin-bottom:16px;">Ruta: <strong>{request.url.path}</strong></p>
                <pre style="background:#1e293b; color:#4ade80; padding:20px; border-radius:8px; overflow-x:auto;">{trace}</pre>
                <p style="margin-top:20px;"><a href="/pedidos" style="color:#fff; background:#2563eb; padding:10px 18px; border-radius:6px; text-decoration:none;">Volver a Órdenes de Compra</a></p>
            </div>
            """,
            status_code=500
        )


app.add_middleware(SessionMiddleware, secret_key=cargar_secret_key(), max_age=8 * 3600, same_site="lax")
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))
templates.env.filters["money"] = lambda v: f"${Decimal(str(v or 0)):,.2f}"


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


class NoAutenticado(Exception):
    pass


@app.exception_handler(NoAutenticado)
async def _redirigir_login(request: Request, exc: NoAutenticado):
    return RedirectResponse("/login", status_code=303)


def usuario_actual(request: Request, db: Session = Depends(get_db)) -> models.Asesor:
    uid = request.session.get("asesor_id")
    user = db.get(models.Asesor, uid) if uid else None
    if not user:
        request.session.clear()
        raise NoAutenticado()
    return user


def ids_asesores_visibles(db: Session, user: models.Asesor) -> list[int] | None:
    if user.rol in ["Administrador", "Gerencia"]:
        return None
    if user.rol == "Jefe de area":
        sub_ids = [a.id for a in db.query(models.Asesor).filter(models.Asesor.jefe_id == user.id).all()]
        sub_ids.append(user.id)
        return sub_ids
    return [user.id]


def tiene_permiso_gestion(user: models.Asesor, objetivo: models.Asesor) -> bool:
    if user.rol in ["Administrador", "Gerencia"]:
        return True
    if user.rol == "Jefe de area" and (objetivo.jefe_id == user.id or objetivo.id == user.id):
        return True
    return False


def render(request: Request, nombre: str, ctx: dict | None = None, status: int = 200):
    ctx = dict(ctx or {})
    ctx.setdefault("flash", request.session.pop("flash", None))
    ctx["request"] = request
    ctx["roles_superiores"] = ROLES_SUPERIORES
    return templates.TemplateResponse(request=request, name=nombre, context=ctx, status_code=status)


EMAIL_RE = re.compile(r"[^@\s]+@[^@\s]+\.[^@\s]+")
RESET_MINUTOS = 30


def config_correo() -> dict:
    cfg = {
        "host": "mail.flylogistic.aero",
        "port": 465,
        "user": CORREO_OFICIAL,
        "password": CLAVE_CORREO_PRUEBA,
        "from": f"Fly Logistic S.A.S. <{CORREO_OFICIAL}>",
        "app_url": "http://127.0.0.1:8000"
    }
    for filename in ("correo_config.json", "correo_config.json.json"):
        ruta = os.path.join(BASE_DIR, filename)
        if os.path.exists(ruta):
            try:
                with open(ruta, encoding="utf-8") as f:
                    leido = json.load(f)
                    cfg.update({k: v for k, v in leido.items() if v is not None and str(v).strip() != ""})
                break
            except Exception:
                pass
    for k, env in (("host", "SMTP_HOST"), ("port", "SMTP_PORT"), ("user", "SMTP_USER"),
                   ("password", "SMTP_PASSWORD"), ("from", "SMTP_FROM"), ("app_url", "APP_URL")):
        if os.getenv(env):
            cfg[k] = os.getenv(env)
    return cfg


def guardar_config_correo(host: str, port: int, user_smtp: str, password_smtp: str):
    cfg = config_correo()
    cfg.update({
        "host": host.strip(),
        "port": int(port),
        "user": user_smtp.strip(),
        "password": password_smtp,
        "from": f"Fly Logistic S.A.S. <{user_smtp.strip()}>"
    })
    for filename in ("correo_config.json", "correo_config.json.json"):
        try:
            with open(os.path.join(BASE_DIR, filename), "w", encoding="utf-8") as f:
                json.dump(cfg, f, indent=2, ensure_ascii=False)
        except Exception:
            pass


def enviar_correo_con_diagnostico(destino: str, asunto: str, cuerpo: str) -> tuple[bool, str]:
    cfg = config_correo()
    remitente = cfg.get("user") or CORREO_OFICIAL
    clave = cfg.get("password") or CLAVE_CORREO_PRUEBA
    host_principal = (cfg.get("host") or "mail.flylogistic.aero").strip()
    puerto_principal = int(cfg.get("port") or 465)

    msg = EmailMessage()
    msg["From"] = cfg.get("from") or f"Fly Logistic S.A.S. <{remitente}>"
    msg["To"] = destino
    msg["Subject"] = asunto
    msg.set_content(cuerpo)

    intentos = [
        (host_principal, puerto_principal),
        ("mail.flylogistic.aero", 465),
        ("mail.flylogistic.aero", 587),
        ("smtp.office365.com", 587),
        ("smtp.gmail.com", 465),
        ("smtp.gmail.com", 587),
    ]
    vistos = set()
    errores = []

    ctx_ssl = ssl.create_default_context()
    ctx_sin_verificar = ssl._create_unverified_context()

    for host, port in intentos:
        if (host, port) in vistos:
            continue
        vistos.add((host, port))
        for contexto in (ctx_ssl, ctx_sin_verificar):
            try:
                if int(port) == 465:
                    with smtplib.SMTP_SSL(host, int(port), context=contexto, timeout=12) as s:
                        s.ehlo()
                        s.login(remitente, clave)
                        s.send_message(msg)
                else:
                    with smtplib.SMTP(host, int(port), timeout=12) as s:
                        s.ehlo()
                        s.starttls(context=contexto)
                        s.ehlo()
                        s.login(remitente, clave)
                        s.send_message(msg)
                detalle_ok = f"Correo entregado vía {host}:{port} a {destino}"
                return True, detalle_ok
            except Exception as e:
                errores.append(f"{host}:{port} -> {e}")

    resumen_err = " | ".join(errores[:3])
    return False, resumen_err


def enviar_correo(destino: str, asunto: str, cuerpo: str) -> bool:
    ok, _ = enviar_correo_con_diagnostico(destino, asunto, cuerpo)
    return ok


def construir_cuerpo_alerta_po(p: models.Pedido) -> tuple[str, str]:
    info_dias = p.dias_habiles_info
    estado_semaforo = "VENCIDO / RETRASADO" if info_dias["en_contra"] else "EN TIEMPO"
    asunto = f"[Fly Logistic S.A.S.] Seguimiento {info_dias['etiqueta']} - Cliente: {p.cliente.nombre}"

    lineas_items = []
    for idx, it in enumerate(p.items, start=1):
        alt_txt = f" (Alterno: {it.parte_numero_alterno})" if it.parte_numero_alterno else ""
        est_prod = it.estado_producto or "PENDIENTE"
        lineas_items.append(
            f"  {idx}. [{est_prod}] Part Number: {it.parte_numero}{alt_txt} | {it.descripcion} | "
            f"Cant: {it.cantidad} {it.unidad_medida} | Unit: ${Decimal(str(it.precio_unitario)):,.2f} | "
            f"Impuesto: {Decimal(str(it.iva_porcentaje)):g}% | Total Línea: ${it.total_item:,.2f}"
        )

    intl_txt = ""
    if p.tipo == "Internacional":
        i = p.internacional
        llega_cm = i.llega_click_mail if i else "No"
        fecha_cm = i.fecha_click_mail.isoformat() if (i and i.fecha_click_mail) else "No aplica"
        fecha_nac = i.fecha_nacionalizacion.isoformat() if (i and i.fecha_nacionalizacion) else "Pendiente"
        est_nac = i.estado_nacionalizacion if (i and i.estado_nacionalizacion) else "Pendiente"
        intl_txt = (
            f"\nDATOS INTERNACIONALES Y NACIONALIZACIÓN\n"
            f"- Recepción Warehouse EE. UU.: {llega_cm} (Fecha: {fecha_cm})\n"
            f"- Tracking: {(i.tracking if i and i.tracking else 'Sin asignar')}\n"
            f"- Estado de Nacionalización: {est_nac}\n"
            f"- Fecha de Nacionalización: {fecha_nac}\n"
            f"- Costo de Nacionalización: ${p.valor_nacionalizacion_dec:,.2f}\n"
        )

    cuerpo = (
        f"Cordial saludo, {p.asesor.nombre}.\n\n"
        f"Se adjunta el resumen operativo de la Orden de Compra {p.po_codigo}:\n\n"
        f"- Orden de Compra: {p.po_codigo}\n"
        f"- Cliente: {p.cliente.nombre}\n"
        f"- Operación: {p.tipo}\n"
        f"- Estado General: {p.estado}\n"
        f"- Fecha de Inicio: {p.fecha_inicial.isoformat()}\n"
        f"- Fecha Prometida: {p.plazo_entrega.isoformat() if p.plazo_entrega else 'Sin definir'}\n"
        f"- Estado de Plazo: {info_dias['etiqueta']} ({estado_semaforo})\n"
        f"- Condición de Pago: {p.metodo_pago or 'Sin definir'}\n"
        f"{intl_txt}\n"
        f"DETALLE DE ÍTEMS\n"
        + "\n".join(lineas_items) +
        f"\n\nTOTALES DE LA ORDEN\n"
        f"- Subtotal: ${p.subtotal:,.2f}\n"
        f"- Impuestos: ${p.total_impuestos:,.2f}\n"
        f"- Nacionalización: ${p.valor_nacionalizacion_dec:,.2f}\n"
        f"- Total General: ${p.total:,.2f}\n\n"
        f"Fly Logistic S.A.S."
    )
    return asunto, cuerpo


CAMPOS = [
    "fecha_inicial", "po_numero", "cliente", "tipo_pedido", "metodo_pago", "iva_porcentaje",
    "valor_nacionalizacion", "plazo_entrega", "fecha_entrega_cliente", "numero_remision",
    "llega_click_mail", "fecha_click_mail", "tracking", "estado_nacionalizacion",
    "fecha_nacionaliza", "documentacion", "estado_general", "observaciones", "asesor_asignado_id"
]
CAMPOS_ITEM = [
    "fecha_compra_prov", "proveedor", "parte_numero", "parte_numero_alterno",
    "unidad_medida", "descripcion", "cantidad", "precio_unitario", "iva_item", "estado_producto"
]


def datos_desde_form(f) -> dict:
    d = {k: (f.get(k) or "").strip() if isinstance(f.get(k), str) else "" for k in CAMPOS}
    metodos_lista = [m.strip() for m in f.getlist("metodos_pago_check") if isinstance(m, str) and m.strip()]
    if metodos_lista:
        desgloses = []
        for m in metodos_lista:
            monto_m = (f.get(f"monto_{m}") or "").strip()
            if monto_m:
                desgloses.append(f"{m}: ${monto_m}")
            else:
                desgloses.append(m)
        d["metodo_pago"] = ", ".join(desgloses)

    listas_raw = {k: [str(v).strip() for v in f.getlist(k)] for k in CAMPOS_ITEM}
    num_filas = max((len(listas_raw["parte_numero"]), len(listas_raw["descripcion"]), len(listas_raw["proveedor"])), default=0)

    items = []
    for idx in range(num_filas):
        fila_dict = {}
        for k in CAMPOS_ITEM:
            vals = listas_raw.get(k, [])
            if idx < len(vals):
                fila_dict[k] = vals[idx]
            else:
                fila_dict[k] = "PENDIENTE" if k == "estado_producto" else ("UND" if k == "unidad_medida" else "")
        if any(fila_dict.get(c) for c in ("proveedor", "parte_numero", "descripcion", "precio_unitario")):
            items.append(fila_dict)
    d["items"] = items
    return d


def datos_desde_pedido(p: models.Pedido) -> dict:
    iso = lambda v: v.isoformat() if v else ""
    i = p.internacional
    llega_cm = i.llega_click_mail if (i and i.llega_click_mail) else ("Si" if (i and i.fecha_click_mail) else "No")
    return {
        "fecha_inicial": iso(p.fecha_inicial),
        "po_numero": p.po,
        "cliente": p.cliente.nombre,
        "tipo_pedido": p.tipo,
        "metodo_pago": p.metodo_pago or "",
        "iva_porcentaje": f"{Decimal(str(p.iva_porcentaje)):g}",
        "valor_nacionalizacion": f"{Decimal(str(p.valor_nacionalizacion or 0)):.2f}",
        "plazo_entrega": iso(p.plazo_entrega),
        "fecha_entrega_cliente": iso(p.fecha_entrega_cliente),
        "numero_remision": p.numero_remision or "",
        "llega_click_mail": llega_cm,
        "fecha_click_mail": iso(i.fecha_click_mail) if i else "",
        "tracking": (i.tracking or "") if i else "",
        "estado_nacionalizacion": (i.estado_nacionalizacion or "") if i else "",
        "fecha_nacionaliza": iso(i.fecha_nacionalizacion) if i else "",
        "documentacion": p.documentacion or "",
        "documentacion_archivo": p.documentacion_archivo or "",
        "documentacion_drive_link": p.documentacion_drive_link or "",
        "estado_general": p.estado,
        "observaciones": p.observaciones or "",
        "asesor_asignado_id": str(p.asesor_id),
        "items": [{
            "fecha_compra_prov": iso(it.fecha_compra_proveedor),
            "proveedor": it.proveedor,
            "parte_numero": it.parte_numero,
            "parte_numero_alterno": it.parte_numero_alterno or "",
            "unidad_medida": it.unidad_medida or "UND",
            "descripcion": it.descripcion,
            "cantidad": str(it.cantidad),
            "precio_unitario": f"{Decimal(str(it.precio_unitario)):.2f}",
            "iva_item": f"{Decimal(str(it.iva_porcentaje if it.iva_porcentaje is not None else p.iva_porcentaje)):g}",
            "estado_producto": it.estado_producto or "PENDIENTE",
        } for it in p.items],
    }


def _fecha(txt, campo, obligatoria=False):
    if not txt:
        if obligatoria:
            raise ValueError(f"La fecha «{campo}» es obligatoria.")
        return None
    try:
        return datetime.strptime(txt, "%Y-%m-%d").date()
    except ValueError:
        raise ValueError(f"Fecha inválida en «{campo}».")


def _decimal(txt, campo):
    try:
        v = Decimal(txt)
    except (InvalidOperation, TypeError):
        raise ValueError(f"Número inválido en «{campo}».")
    if not v.is_finite():
        raise ValueError(f"Número inválido en «{campo}».")
    return v


def parsear(d: dict, tipo_fijo: str | None = None) -> dict:
    po_limpio = (d.get("po_numero") or "").strip()
    if not po_limpio:
        raise ValueError("El número de PO (Orden de Compra) es obligatorio.")
    if not (d.get("cliente") or "").strip():
        raise ValueError("El nombre del cliente es obligatorio.")
    tipo = tipo_fijo or d.get("tipo_pedido")
    if tipo not in ("Nacional", "Internacional"):
        raise ValueError("Seleccione el tipo de operación (Nacional o Internacional).")

    fecha_ini = _fecha(d.get("fecha_inicial"), "Fecha inicial", obligatoria=True)
    plazo_ent = _fecha(d.get("plazo_entrega"), "Plazo de entrega prometido", obligatoria=True)
    if not (d.get("metodo_pago") or "").strip():
        raise ValueError("Seleccione al menos un método de pago.")

    iva_defecto = "19" if tipo == "Nacional" else "7"
    iva_cabecera = _decimal(d.get("iva_porcentaje") or iva_defecto, "IVA / TAX")
    if not 0 <= iva_cabecera <= 100:
        raise ValueError("El porcentaje de impuesto debe estar entre 0 y 100.")

    val_nac = _decimal(d.get("valor_nacionalizacion") or "0", "Valor de Nacionalización")
    if val_nac < 0:
        raise ValueError("El valor de nacionalización no puede ser negativo.")

    if d.get("estado_general") not in ESTADOS:
        raise ValueError("Estado general inválido (En proceso o Entregado).")
    if not d.get("items"):
        raise ValueError("Registre al menos un producto en la orden.")

    items = []
    for n, it in enumerate(d["items"], start=1):
        for k, etiqueta in (("proveedor", "proveedor"), ("parte_numero", "Part Number"), ("descripcion", "descripción")):
            if not (it.get(k) or "").strip():
                raise ValueError(f"Ítem {n}: falta completar el campo «{etiqueta}».")
        try:
            cant = int(it.get("cantidad", 0))
        except ValueError:
            raise ValueError(f"Ítem {n}: cantidad inválida.")
        if cant < 1:
            raise ValueError(f"Ítem {n}: la cantidad mínima es 1.")
        if (it.get("precio_unitario") or "").strip() == "":
            raise ValueError(f"Ítem {n}: el precio unitario es obligatorio.")
        precio = _decimal(it.get("precio_unitario"), f"precio unitario del ítem {n}")
        if precio <= 0:
            raise ValueError(f"Ítem {n}: el precio unitario debe ser mayor a 0.")

        iva_it_txt = it.get("iva_item")
        iva_it = _decimal(iva_it_txt if iva_it_txt != "" else str(iva_cabecera), f"impuesto del ítem {n}")
        if not 0 <= iva_it <= 100:
            raise ValueError(f"Ítem {n}: el impuesto debe estar entre 0 y 100.")

        est_prod = (it.get("estado_producto") or "PENDIENTE").strip().upper()
        if est_prod not in ESTADOS_PRODUCTO:
            est_prod = "PENDIENTE"

        items.append({
            "fecha_compra_proveedor": _fecha(it.get("fecha_compra_prov"), f"compra a proveedor (ítem {n})"),
            "proveedor": it["proveedor"].strip(),
            "parte_numero": it["parte_numero"].strip(),
            "parte_numero_alterno": (it.get("parte_numero_alterno") or "").strip(),
            "unidad_medida": it.get("unidad_medida") or "UND",
            "descripcion": it["descripcion"].strip(),
            "cantidad": cant,
            "precio_unitario": precio,
            "iva_porcentaje": iva_it,
            "estado_producto": est_prod,
        })

    if tipo == "Internacional":
        fecha_cm = _fecha(d.get("fecha_click_mail"), "Fecha Warehouse EE. UU.")
        llega_cm = "Si" if (d.get("llega_click_mail") == "Si" or fecha_cm is not None) else "No"
        intl_data = {
            "llega_click_mail": llega_cm,
            "fecha_click_mail": fecha_cm,
            "tracking": (d.get("tracking") or "").strip(),
            "estado_nacionalizacion": d.get("estado_nacionalizacion") or "No aplica / Pendiente",
            "fecha_nacionalizacion": _fecha(d.get("fecha_nacionaliza"), "Fecha de Nacionalización"),
        }
    else:
        intl_data = None

    return {
        "po": po_limpio,
        "tipo": tipo,
        "fecha_inicial": fecha_ini,
        "cliente": d["cliente"].strip(),
        "metodo_pago": d.get("metodo_pago"),
        "iva_porcentaje": iva_cabecera,
        "valor_nacionalizacion": val_nac if tipo == "Internacional" else Decimal("0"),
        "plazo_entrega": plazo_ent,
        "fecha_entrega_cliente": _fecha(d.get("fecha_entrega_cliente"), "Entrega al cliente"),
        "numero_remision": (d.get("numero_remision") or "").strip(),
        "documentacion": (d.get("documentacion") or "").strip(),
        "estado": d["estado_general"],
        "observaciones": (d.get("observaciones") or "").strip(),
        "items": items,
        "intl": intl_data,
    }


def obtener_o_crear_cliente(db: Session, nombre_raw: str) -> models.Cliente:
    nombre_limpio = (nombre_raw or "").strip()
    cli = db.query(models.Cliente).filter(models.Cliente.nombre == nombre_limpio).first()
    if cli:
        return cli
    norm_buscado = models.normalizar(nombre_limpio)
    for c in db.query(models.Cliente).all():
        if models.normalizar(c.nombre) == norm_buscado:
            return c
    cli = models.Cliente(nombre=nombre_limpio)
    db.add(cli)
    db.flush()
    return cli


def aplicar(db: Session, pedido: models.Pedido, v: dict):
    cli = obtener_o_crear_cliente(db, v["cliente"])
    pedido.cliente_id = cli.id
    for k in ("fecha_inicial", "metodo_pago", "iva_porcentaje", "valor_nacionalizacion", "plazo_entrega",
              "fecha_entrega_cliente", "numero_remision", "documentacion", "estado", "observaciones"):
        setattr(pedido, k, v[k])
    pedido.actualizado_en = datetime.utcnow()
    pedido.items.clear()
    for it in v["items"]:
        pedido.items.append(models.ItemPedido(**it))

    if v["intl"]:
        if pedido.internacional:
            for k, val in v["intl"].items():
                setattr(pedido.internacional, k, val)
        else:
            pedido.internacional = models.PedidoInternacional(**v["intl"])
    else:
        pedido.internacional = None


def pedido_del_usuario(db: Session, pid: int, user: models.Asesor) -> models.Pedido | None:
    visibles = ids_asesores_visibles(db, user)
    q = db.query(models.Pedido).options(
        joinedload(models.Pedido.items),
        joinedload(models.Pedido.cliente),
        joinedload(models.Pedido.asesor),
        joinedload(models.Pedido.internacional)
    ).filter(models.Pedido.id == pid)
    if visibles is not None:
        q = q.filter(models.Pedido.asesor_id.in_(visibles))
    return q.first()


@app.api_route("/healthz", methods=["GET", "HEAD"])
def health_check():
    return JSONResponse({"status": "ok", "timestamp": datetime.utcnow().isoformat()})


@app.api_route("/", methods=["GET", "HEAD"])
def inicio(request: Request):
    if request.method == "HEAD":
        return HTMLResponse(content="", status_code=200)
    return RedirectResponse("/pedidos" if request.session.get("asesor_id") else "/login", status_code=303)


@app.api_route("/login", methods=["GET", "HEAD"], response_class=HTMLResponse)
def login_form(request: Request):
    if request.method == "HEAD":
        return HTMLResponse(content="", status_code=200)
    if request.session.get("asesor_id"):
        return RedirectResponse("/pedidos", status_code=303)
    return render(request, "login.html", {"username": "", "error": None})


@app.post("/login", response_class=HTMLResponse)
async def login_post(request: Request, db: Session = Depends(get_db)):
    f = await request.form()
    raw_user = (f.get("username") or "").strip()
    raw_pass = (f.get("password") or "").strip()
    username_norm = models.normalizar(raw_user)
    user = db.query(models.Asesor).filter(
        or_(models.Asesor.username == username_norm, models.Asesor.nombre_norm == username_norm)
    ).first() if username_norm else None

    if not user or not verificar_clave(raw_pass, user.password_hash):
        return render(request, "login.html", {"username": raw_user, "error": "Credenciales de acceso inválidas."}, 401)

    request.session.clear()
    request.session["asesor_id"] = user.id
    return RedirectResponse("/pedidos", status_code=303)


@app.api_route("/logout", methods=["GET", "POST"])
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


@app.get("/registro", response_class=HTMLResponse)
def registro_form(request: Request, db: Session = Depends(get_db)):
    superiores = db.query(models.Asesor).filter(models.Asesor.rol.in_(ROLES_SUPERIORES)).order_by(models.Asesor.nombre).all()
    return render(request, "registro.html", {"v": {"email": ""}, "superiores": superiores, "error": None})


@app.post("/registro", response_class=HTMLResponse)
async def registro_post(request: Request, db: Session = Depends(get_db)):
    f = await request.form()
    v = {k: (f.get(k) or "").strip() for k in ("nombre", "username", "email", "jefe_id")}
    if not v["email"]:
        v["email"] = CORREO_OFICIAL
    usuario = models.normalizar(v["username"])
    password = (f.get("password") or "").strip()
    password2 = (f.get("password2") or "").strip()
    clave_autorizacion = (f.get("clave_maestra") or "").strip()
    jefe_id_int = int(v["jefe_id"]) if v["jefe_id"].isdigit() else None
    superiores = db.query(models.Asesor).filter(models.Asesor.rol.in_(ROLES_SUPERIORES)).order_by(models.Asesor.nombre).all()

    def error(msg, status=400):
        return render(request, "registro.html", {"v": v, "superiores": superiores, "error": msg}, status)

    autorizado, sup_encontrado = buscar_autorizador_por_clave(db, clave_autorizacion, jefe_id_int)
    if not autorizado:
        return error("La clave de autorización ingresada no es válida.", 403)
    if len(v["nombre"]) < 3:
        return error("Ingrese el nombre y apellido completo.")
    if not v["username"]:
        return error("El nombre de usuario es obligatorio.")
    if not EMAIL_RE.fullmatch(v["email"]):
        return error("Ingrese un correo electrónico corporativo válido.")
    if not 3 <= len(usuario) <= 100:
        return error("El usuario debe contener entre 3 y 100 caracteres.")
    if len(password) < 8:
        return error("La contraseña debe tener mínimo 8 caracteres.")
    if password != password2:
        return error("Las contraseñas no coinciden.")
    if db.query(models.Asesor).filter_by(username=usuario).first():
        return error("El nombre de usuario ya se encuentra registrado.", 409)

    asesor = db.query(models.Asesor).filter_by(nombre_norm=models.normalizar(v["nombre"])).first()
    if asesor and asesor.password_hash:
        return error("Ya existe un perfil activo para este colaborador.", 409)
    if not asesor:
        asesor = models.Asesor(nombre=v["nombre"], nombre_norm=models.normalizar(v["nombre"]), rol="Asesor")
        db.add(asesor)

    asesor.username = usuario
    asesor.password_hash = hash_clave(password)
    asesor.email = v["email"]
    asesor.rol = "Asesor"
    asesor.jefe_id = jefe_id_int or (sup_encontrado.id if sup_encontrado else None)
    db.commit()

    request.session.clear()
    request.session["asesor_id"] = asesor.id
    request.session["flash"] = {"ok": True, "msg": f"Perfil corporativo habilitado para {asesor.nombre}."}
    return RedirectResponse("/pedidos", status_code=303)


def _hash_token(token: str) -> str:
    return hashlib.sha256(token.encode()).hexdigest()


def asesor_por_token(db: Session, token: str):
    a = db.query(models.Asesor).filter_by(reset_token_hash=_hash_token(token)).first()
    if a and a.reset_expira and a.reset_expira > datetime.utcnow():
        return a
    return None


@app.get("/olvide", response_class=HTMLResponse)
def olvide_form(request: Request):
    return render(request, "olvide.html", {
        "enviado": False, "valor": "", "bloqueado_asesor": False,
        "error": None, "enlace_directo_prueba": None
    })


@app.post("/olvide", response_class=HTMLResponse)
async def olvide_post(request: Request, db: Session = Depends(get_db)):
    valor = ((await request.form()).get("identificador") or "").strip()
    if not valor:
        return render(request, "olvide.html", {
            "enviado": False, "valor": "", "bloqueado_asesor": False,
            "error": "Ingrese su usuario o correo electrónico.", "enlace_directo_prueba": None
        })

    cuentas = db.query(models.Asesor).filter(
        or_(models.Asesor.username == models.normalizar(valor),
            func.lower(models.Asesor.email) == valor.lower())
    ).all()

    cuentas_superiores = [a for a in cuentas if a.rol in ROLES_SUPERIORES]
    if cuentas and not cuentas_superiores:
        jefe_nombre = cuentas[0].jefe.nombre if cuentas[0].jefe else "Jefatura de Área o Administración"
        return render(request, "olvide.html", {
            "enviado": False,
            "valor": valor,
            "bloqueado_asesor": True,
            "jefe_nombre": jefe_nombre,
            "error": None,
            "enlace_directo_prueba": None
        }, 403)

    ahora = datetime.utcnow()
    base = (config_correo().get("app_url") or str(request.base_url)).rstrip("/")
    enlace_generado = None
    estado_smtp = ""
    for a in cuentas_superiores:
        destino_correo = a.email or CORREO_OFICIAL
        token = secrets.token_urlsafe(32)
        a.reset_token_hash = _hash_token(token)
        a.reset_expira = ahora + timedelta(minutes=RESET_MINUTOS)
        db.commit()
        enlace_generado = f"{base}/restablecer/{token}"
        ok, detalle = enviar_correo_con_diagnostico(
            destino_correo,
            "Restablecimiento de Credenciales - Fly Logistic S.A.S.",
            f"Hola {a.nombre},\n\nSe ha solicitado el restablecimiento de su contraseña.\n"
            f"Enlace de acceso (válido por {RESET_MINUTOS} minutos):\n\n"
            f"{enlace_generado}\n\n"
            f"Fly Logistic S.A.S."
        )
        estado_smtp = detalle

    return render(request, "olvide.html", {
        "enviado": True, "valor": valor, "bloqueado_asesor": False,
        "error": None, "enlace_directo_prueba": enlace_generado, "estado_smtp": estado_smtp
    })


@app.get("/restablecer/{token}", response_class=HTMLResponse)
def restablecer_form(token: str, request: Request, db: Session = Depends(get_db)):
    valido = asesor_por_token(db, token) is not None
    return render(request, "restablecer.html", {"valido": valido, "error": None}, 200 if valido else 400)


@app.post("/restablecer/{token}", response_class=HTMLResponse)
async def restablecer_post(token: str, request: Request, db: Session = Depends(get_db)):
    a = asesor_por_token(db, token)
    if not a:
        return render(request, "restablecer.html", {"valido": False, "error": None}, 400)
    f = await request.form()
    p1, p2 = (f.get("password") or "").strip(), (f.get("password2") or "").strip()
    if len(p1) < 8:
        return render(request, "restablecer.html", {"valido": True, "error": "Mínimo 8 caracteres."}, 400)
    if p1 != p2:
        return render(request, "restablecer.html", {"valido": True, "error": "Las contraseñas no coinciden."}, 400)
    a.password_hash = hash_clave(p1)
    a.reset_token_hash = None
    a.reset_expira = None
    db.commit()
    request.session.clear()
    request.session["flash"] = {"ok": True, "msg": "Contraseña actualizada correctamente."}
    return RedirectResponse("/login", status_code=303)


@app.get("/perfil", response_class=HTMLResponse)
def perfil(request: Request, user: models.Asesor = Depends(usuario_actual)):
    return render(request, "perfil.html", {"user": user, "cfg_smtp": config_correo(), "error": None})


@app.post("/perfil/correo", response_class=HTMLResponse)
async def perfil_correo(request: Request, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    email = ((await request.form()).get("email") or "").strip()
    if not EMAIL_RE.fullmatch(email):
        return render(request, "perfil.html", {"user": user, "cfg_smtp": config_correo(), "error": "Formato de correo inválido."}, 400)
    user.email = email
    db.commit()
    request.session["flash"] = {"ok": True, "msg": "Correo electrónico actualizado."}
    return RedirectResponse("/perfil", status_code=303)


@app.post("/perfil/datos", response_class=HTMLResponse)
async def perfil_datos(request: Request, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    f = await request.form()
    email = (f.get("email") or "").strip()
    telefono = (f.get("telefono") or "").strip()
    if not EMAIL_RE.fullmatch(email):
        return render(request, "perfil.html", {"user": user, "cfg_smtp": config_correo(), "error": "Formato de correo inválido."}, 400)
    user.email = email
    user.telefono = telefono
    db.commit()
    request.session["flash"] = {"ok": True, "msg": "Datos de contacto actualizados."}
    return RedirectResponse("/perfil", status_code=303)


@app.post("/perfil/smtp", response_class=HTMLResponse)
async def perfil_smtp(request: Request, user: models.Asesor = Depends(usuario_actual)):
    if user.rol not in ROLES_SUPERIORES:
        return RedirectResponse("/perfil", status_code=303)
    f = await request.form()
    host = (f.get("smtp_host") or "mail.flylogistic.aero").strip()
    port = int(f.get("smtp_port") or 465)
    smtp_user = (f.get("smtp_user") or CORREO_OFICIAL).strip()
    smtp_pass = (f.get("smtp_pass") or CLAVE_CORREO_PRUEBA).strip()
    destino_prueba = (f.get("destino_prueba") or user.email or CORREO_OFICIAL).strip()

    guardar_config_correo(host, port, smtp_user, smtp_pass)

    if f.get("accion_smtp") == "probar":
        ok, detalle = enviar_correo_con_diagnostico(
            destino_prueba,
            "Prueba de Servicio SMTP - Fly Logistic S.A.S.",
            f"Hola {user.nombre},\n\nConfirmación de salida de correo desde el sistema OMS.\n"
            f"Servidor: {host}:{port}\nCuenta: {smtp_user}\n"
        )
        if ok:
            request.session["flash"] = {"ok": True, "msg": f"Conexión SMTP verificada: {detalle}"}
        else:
            request.session["flash"] = {"ok": False, "msg": f"Parámetros guardados. Respuesta del servidor: {detalle}"}
    else:
        request.session["flash"] = {"ok": True, "msg": "Configuración SMTP actualizada."}
    return RedirectResponse("/perfil", status_code=303)


@app.post("/perfil/clave", response_class=HTMLResponse)
async def perfil_clave(request: Request, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    f = await request.form()
    actual, p1, p2 = (f.get("actual") or "").strip(), (f.get("password") or "").strip(), (f.get("password2") or "").strip()
    err = None
    if not verificar_clave(actual, user.password_hash):
        err = "La contraseña actual es incorrecta."
    elif len(p1) < 8:
        err = "La nueva contraseña debe tener mínimo 8 caracteres."
    elif p1 != p2:
        err = "Las contraseñas nuevas no coinciden."
    if err:
        return render(request, "perfil.html", {"user": user, "cfg_smtp": config_correo(), "error": err}, 400)
    user.password_hash = hash_clave(p1)
    user.reset_token_hash, user.reset_expira = None, None
    db.commit()
    request.session["flash"] = {"ok": True, "msg": "Contraseña actualizada."}
    return RedirectResponse("/perfil", status_code=303)


@app.get("/gestion_usuarios", response_class=HTMLResponse)
def gestion_usuarios(request: Request, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    if user.rol not in ROLES_ACCESO_USUARIOS:
        return RedirectResponse("/pedidos", status_code=303)

    if user.rol in ["Administrador", "Gerencia"]:
        usuarios = db.query(models.Asesor).order_by(models.Asesor.id).all()
        jefes_disponibles = db.query(models.Asesor).filter(models.Asesor.rol.in_(ROLES_SUPERIORES)).all()
    else:
        usuarios = db.query(models.Asesor).filter(
            or_(models.Asesor.jefe_id == user.id, models.Asesor.id == user.id)
        ).order_by(models.Asesor.id).all()
        jefes_disponibles = [user]

    return render(request, "Usuarios.html", {
        "user": user,
        "usuarios": usuarios,
        "jefes_disponibles": jefes_disponibles,
        "roles_sistema": ROLES_SISTEMA
    })


@app.post("/gestion_usuarios/accion", response_class=HTMLResponse)
async def accion_usuario(request: Request, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    if user.rol not in ROLES_ACCESO_USUARIOS:
        return RedirectResponse("/pedidos", status_code=303)

    f = await request.form()
    accion = (f.get("accion") or "").strip()
    objetivo_id = (f.get("usuario_id") or "").strip()
    objetivo = db.get(models.Asesor, int(objetivo_id)) if objetivo_id.isdigit() else None

    if accion == "crear":
        nombre = (f.get("nombre") or "").strip()
        username = models.normalizar(f.get("username") or "")
        email = (f.get("email") or "").strip() or CORREO_OFICIAL
        password = (f.get("password") or "").strip() or CLAVE_CORREO_PRUEBA
        rol_elegido = (f.get("rol") or "Asesor").strip()
        jefe_id_str = (f.get("jefe_id") or "").strip()

        if len(nombre) < 3 or len(username) < 3 or len(password) < 8:
            request.session["flash"] = {"ok": False, "msg": "Verifique nombre, usuario (mín. 3 caracteres) y contraseña (mín. 8 caracteres)."}
            return RedirectResponse("/gestion_usuarios", status_code=303)

        if db.query(models.Asesor).filter(
            or_(models.Asesor.username == username, models.Asesor.nombre_norm == models.normalizar(nombre))
        ).first():
            request.session["flash"] = {"ok": False, "msg": "Ya existe un colaborador registrado con ese nombre o usuario."}
            return RedirectResponse("/gestion_usuarios", status_code=303)

        rol_final = rol_elegido if (user.rol in ["Administrador", "Gerencia"] and rol_elegido in ROLES_SISTEMA) else "Asesor"
        if user.rol == "Jefe de area":
            jefe_final = user.id
        else:
            jefe_final = int(jefe_id_str) if jefe_id_str.isdigit() else None

        nuevo_user = models.Asesor(
            nombre=nombre,
            nombre_norm=models.normalizar(nombre),
            username=username,
            email=email,
            rol=rol_final,
            jefe_id=jefe_final,
            password_hash=hash_clave(password)
        )
        db.add(nuevo_user)
        db.commit()
        request.session["flash"] = {"ok": True, "msg": f"Colaborador {nuevo_user.nombre} ({nuevo_user.rol}) registrado."}
        return RedirectResponse("/gestion_usuarios", status_code=303)

    if not objetivo or not tiene_permiso_gestion(user, objetivo):
        request.session["flash"] = {"ok": False, "msg": "Permisos insuficientes para modificar este perfil."}
        return RedirectResponse("/gestion_usuarios", status_code=303)

    if accion == "modificar_clave":
        nueva_clave = (f.get("nueva_clave") or "").strip()
        if len(nueva_clave) < 8:
            request.session["flash"] = {"ok": False, "msg": "La contraseña debe tener al menos 8 caracteres."}
        else:
            objetivo.password_hash = hash_clave(nueva_clave)
            db.commit()
            request.session["flash"] = {"ok": True, "msg": f"Contraseña actualizada para {objetivo.nombre}."}

    elif accion == "eliminar":
        if objetivo.id == user.id or objetivo.username == "admin":
            request.session["flash"] = {"ok": False, "msg": "La cuenta administrativa principal no puede ser eliminada."}
            return RedirectResponse("/gestion_usuarios", status_code=303)
        try:
            nombre_borrado = objetivo.nombre
            uid_borrar = objetivo.id
            db.query(models.Pedido).filter(models.Pedido.asesor_id == uid_borrar).update(
                {"asesor_id": user.id}, synchronize_session=False
            )
            db.query(models.Asesor).filter(models.Asesor.jefe_id == uid_borrar).update(
                {"jefe_id": user.id}, synchronize_session=False
            )
            db.execute(text("DELETE FROM asesores WHERE id = :uid"), {"uid": uid_borrar})
            db.commit()
            request.session["flash"] = {"ok": True, "msg": f"Usuario {nombre_borrado} dado de baja. Sus órdenes fueron reasignadas."}
        except Exception as e:
            db.rollback()
            request.session["flash"] = {"ok": False, "msg": f"Error al eliminar cuenta: {e}"}

    elif accion == "modificar_correo":
        nuevo_correo = (f.get("nuevo_correo") or "").strip()
        if not EMAIL_RE.fullmatch(nuevo_correo):
            request.session["flash"] = {"ok": False, "msg": "Formato de correo electrónico inválido."}
        else:
            objetivo.email = nuevo_correo
            db.commit()
            request.session["flash"] = {"ok": True, "msg": f"Correo actualizado para {objetivo.nombre}."}

    elif accion == "modificar_rol_jefe":
        nuevo_rol = (f.get("nuevo_rol") or "").strip()
        nuevo_jefe_str = (f.get("nuevo_jefe_id") or "").strip()
        if user.rol in ["Administrador", "Gerencia"] and nuevo_rol in ROLES_SISTEMA:
            objetivo.rol = nuevo_rol
            objetivo.jefe_id = int(nuevo_jefe_str) if nuevo_jefe_str.isdigit() else None
            db.commit()
            request.session["flash"] = {"ok": True, "msg": f"Cargo y asignación actualizados para {objetivo.nombre}."}
        elif user.rol == "Jefe de area" and nuevo_rol in ["Asesor", "Jefe de area"]:
            objetivo.rol = nuevo_rol
            db.commit()
            request.session["flash"] = {"ok": True, "msg": f"Cargo actualizado para {objetivo.nombre}."}

    return RedirectResponse("/gestion_usuarios", status_code=303)


@app.get("/supervision", response_class=HTMLResponse)
def vista_supervision(request: Request, asesor_id: str = "", q: str = "",
                      user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    if user.rol not in ROLES_SUPERIORES:
        return RedirectResponse("/pedidos", status_code=303)

    visibles = ids_asesores_visibles(db, user)
    q_asesores = db.query(models.Asesor).options(
        joinedload(models.Asesor.pedidos).joinedload(models.Pedido.cliente),
        joinedload(models.Asesor.pedidos).joinedload(models.Pedido.items)
    )
    if visibles is not None:
        q_asesores = q_asesores.filter(models.Asesor.id.in_(visibles))
    asesores_lista = q_asesores.order_by(models.Asesor.nombre).all()

    q_pedidos = db.query(models.Pedido).options(
        joinedload(models.Pedido.cliente),
        joinedload(models.Pedido.asesor),
        joinedload(models.Pedido.items),
        joinedload(models.Pedido.internacional)
    )
    if visibles is not None:
        q_pedidos = q_pedidos.filter(models.Pedido.asesor_id.in_(visibles))
    if asesor_id.isdigit():
        q_pedidos = q_pedidos.filter(models.Pedido.asesor_id == int(asesor_id))
    if q.strip():
        like = f"%{q.strip()}%"
        q_pedidos = q_pedidos.join(models.Cliente).filter(
            models.Pedido.po.ilike(like) | models.Cliente.nombre.ilike(like)
        )

    pedidos_supervision = q_pedidos.order_by(models.Pedido.id.desc()).all()

    resumen_asesores = []
    for a in asesores_lista:
        pos_activas = [p for p in a.pedidos if p.estado != "Entregado"]
        resumen_asesores.append({
            "asesor": a,
            "total_pos": len(a.pedidos),
            "pos_activas": len(pos_activas),
            "pedidos": sorted(a.pedidos, key=lambda x: x.dias_habiles_info["orden_urgencia"])
        })

    return render(request, "Supervision.html", {
        "user": user,
        "asesores_lista": asesores_lista,
        "resumen_asesores": resumen_asesores,
        "pedidos_supervision": pedidos_supervision,
        "f_asesor": asesor_id,
        "q": q
    })


@app.post("/supervision/reasignar")
async def reasignar_po_supervision(request: Request, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    if user.rol not in ROLES_SUPERIORES:
        return RedirectResponse("/pedidos", status_code=303)

    f = await request.form()
    modo_transferencia = (f.get("modo") or "individual").strip()
    nuevo_asesor_id = (f.get("nuevo_asesor_id") or "").strip()

    if not nuevo_asesor_id.isdigit():
        request.session["flash"] = {"ok": False, "msg": "Seleccione un asesor destino válido."}
        return RedirectResponse("/supervision", status_code=303)

    asesor_destino = db.get(models.Asesor, int(nuevo_asesor_id))
    if not asesor_destino:
        request.session["flash"] = {"ok": False, "msg": "El asesor destino no existe."}
        return RedirectResponse("/supervision", status_code=303)

    visibles = ids_asesores_visibles(db, user)
    if visibles is not None and asesor_destino.id not in visibles:
        request.session["flash"] = {"ok": False, "msg": "No tiene permisos sobre el asesor destino."}
        return RedirectResponse("/supervision", status_code=303)

    if modo_transferencia == "masivo":
        origen_id = (f.get("asesor_origen_id") or "").strip()
        if not origen_id.isdigit() or int(origen_id) == asesor_destino.id:
            request.session["flash"] = {"ok": False, "msg": "Seleccione un asesor origen diferente al destino."}
            return RedirectResponse("/supervision", status_code=303)
        asesor_origen = db.get(models.Asesor, int(origen_id))
        if not asesor_origen or (visibles is not None and asesor_origen.id not in visibles):
            request.session["flash"] = {"ok": False, "msg": "Sin permisos sobre el asesor origen."}
            return RedirectResponse("/supervision", status_code=303)

        pedidos_mover = db.query(models.Pedido).options(
            joinedload(models.Pedido.items),
            joinedload(models.Pedido.cliente),
            joinedload(models.Pedido.asesor),
            joinedload(models.Pedido.internacional)
        ).filter(models.Pedido.asesor_id == asesor_origen.id).all()

        cantidad_movida = len(pedidos_mover)
        for p in pedidos_mover:
            p.asesor_id = asesor_destino.id
            p.actualizado_en = datetime.utcnow()
        db.commit()

        for p in pedidos_mover:
            db.refresh(p)
            generar_salidas(db, p)

        request.session["flash"] = {
            "ok": True,
            "msg": f"Se transfirieron {cantidad_movida} órdenes de {asesor_origen.nombre} a {asesor_destino.nombre}."
        }
        return RedirectResponse("/supervision", status_code=303)

    pedido_id = (f.get("pedido_id") or "").strip()
    if not pedido_id.isdigit():
        request.session["flash"] = {"ok": False, "msg": "Seleccione una orden válida."}
        return RedirectResponse("/supervision", status_code=303)

    pedido = pedido_del_usuario(db, int(pedido_id), user)
    if not pedido:
        request.session["flash"] = {"ok": False, "msg": "Orden no encontrada o sin permisos."}
        return RedirectResponse("/supervision", status_code=303)

    asesor_anterior = pedido.asesor.nombre if pedido.asesor else "Anterior"
    pedido.asesor_id = asesor_destino.id
    pedido.actualizado_en = datetime.utcnow()
    db.commit()
    db.refresh(pedido)
    generar_salidas(db, pedido)

    request.session["flash"] = {
        "ok": True,
        "msg": f"Orden {pedido.po_codigo} reasignada de {asesor_anterior} a {asesor_destino.nombre}."
    }
    return RedirectResponse("/supervision", status_code=303)


def _construir_analitica(db: Session, user: models.Asesor, asesor_filtro: str = "", tipo_filtro: str = "", desde: str = "", hasta: str = ""):
    visibles = ids_asesores_visibles(db, user)
    q_asesores = db.query(models.Asesor)
    if visibles is not None:
        q_asesores = q_asesores.filter(models.Asesor.id.in_(visibles))
    asesores_equipo = q_asesores.order_by(models.Asesor.nombre).all()

    q_pedidos = db.query(models.Pedido).options(
        joinedload(models.Pedido.items),
        joinedload(models.Pedido.cliente),
        joinedload(models.Pedido.asesor),
        joinedload(models.Pedido.internacional)
    )
    if visibles is not None:
        q_pedidos = q_pedidos.filter(models.Pedido.asesor_id.in_(visibles))

    if asesor_filtro.isdigit():
        q_pedidos = q_pedidos.filter(models.Pedido.asesor_id == int(asesor_filtro))
    if tipo_filtro in ("Nacional", "Internacional"):
        q_pedidos = q_pedidos.filter(models.Pedido.tipo == tipo_filtro)
    if desde:
        try:
            q_pedidos = q_pedidos.filter(models.Pedido.fecha_inicial >= datetime.strptime(desde, "%Y-%m-%d").date())
        except ValueError:
            pass
    if hasta:
        try:
            q_pedidos = q_pedidos.filter(models.Pedido.fecha_inicial <= datetime.strptime(hasta, "%Y-%m-%d").date())
        except ValueError:
            pass

    pedidos = q_pedidos.order_by(models.Pedido.fecha_inicial.desc(), models.Pedido.id.desc()).all()
    hoy = date.today()

    total_pedidos = len(pedidos)
    total_nacionales = sum(1 for p in pedidos if p.tipo == "Nacional")
    total_internacionales = sum(1 for p in pedidos if p.tipo == "Internacional")
    valor_total_ventas = sum(float(p.total) for p in pedidos)
    valor_subtotal = sum(float(p.subtotal) for p in pedidos)
    valor_total_nacionalizacion = sum(float(p.valor_nacionalizacion_dec) for p in pedidos)

    suma_probabilidad = 0
    tiempos_gestion = []
    entregas_a_tiempo = 0
    entregas_evaluadas = 0
    compras_pendientes_po = []
    todos_items_global = []
    items_sin_compra = 0

    ventas_mensuales_nac = {}
    ventas_mensuales_int = {}
    estado_counts = {e: 0 for e in ESTADOS}
    proveedores_top = {}

    metricas_por_asesor = {}
    for a in asesores_equipo:
        metricas_por_asesor[a.id] = {
            "id": a.id,
            "nombre": a.nombre,
            "rol": a.rol,
            "jefe": a.jefe.nombre if a.jefe else "Directo",
            "total_pos": 0,
            "entregados": 0,
            "pendientes": 0,
            "items_pendientes_compra": 0,
            "volumen": 0.0,
            "nacionalizacion": 0.0,
            "dias_acum": 0,
            "dias_count": 0,
            "prob_acum": 0,
            "sla_ok": 0,
            "sla_total": 0,
        }

    for p in pedidos:
        prob_p = PESOS_PROBABILIDAD.get(p.estado, 40)
        suma_probabilidad += prob_p
        estado_counts[p.estado] = estado_counts.get(p.estado, 0) + 1

        mes_key = p.fecha_inicial.strftime("%Y-%m")
        if p.tipo == "Nacional":
            ventas_mensuales_nac[mes_key] = ventas_mensuales_nac.get(mes_key, 0.0) + float(p.total)
        else:
            ventas_mensuales_int[mes_key] = ventas_mensuales_int.get(mes_key, 0.0) + float(p.total)

        cli_nom = p.cliente.nombre if p.cliente else "Sin Cliente"
        fecha_fin = p.fecha_entrega_cliente or hoy
        dias_gest = max(0, (fecha_fin - p.fecha_inicial).days)
        tiempos_gestion.append(dias_gest)

        info_habiles = p.dias_habiles_info
        en_tiempo = not info_habiles["en_contra"]
        if p.plazo_entrega:
            entregas_evaluadas += 1
            if en_tiempo:
                entregas_a_tiempo += 1

        faltan_compras_item = 0
        for it in p.items:
            prov_nom = it.proveedor or "Sin Proveedor"
            proveedores_top[prov_nom] = proveedores_top.get(prov_nom, 0.0) + float(it.subtotal_item)
            if not it.fecha_compra_proveedor or it.estado_producto == "PENDIENTE":
                faltan_compras_item += 1
                items_sin_compra += 1

            todos_items_global.append({
                "pedido_id": p.id,
                "po": p.po_codigo,
                "tipo": p.tipo,
                "asesor": p.asesor.nombre if p.asesor else "",
                "cliente": cli_nom,
                "proveedor": prov_nom,
                "fecha_compra": it.fecha_compra_proveedor.isoformat() if it.fecha_compra_proveedor else "FALTA COMPRAR",
                "comprado": bool(it.fecha_compra_proveedor),
                "parte_numero": it.parte_numero,
                "descripcion": it.descripcion,
                "cantidad": it.cantidad,
                "unidad_medida": it.unidad_medida or "UND",
                "estado_producto": it.estado_producto or "PENDIENTE",
                "estado_po": p.estado,
                "semaforo": info_habiles["etiqueta"],
            })

        if p.estado != "Entregado" or faltan_compras_item > 0:
            compras_pendientes_po.append({
                "id": p.id,
                "po": p.po_codigo,
                "cliente": cli_nom,
                "asesor": p.asesor.nombre if p.asesor else "",
                "asesor_email": (p.asesor.email if p.asesor and p.asesor.email else CORREO_OFICIAL),
                "tipo": p.tipo,
                "estado": p.estado,
                "dias": dias_gest,
                "plazo": p.plazo_entrega.isoformat() if p.plazo_entrega else "Sin definir",
                "estado_plazo": info_habiles["etiqueta"],
                "dias_habiles_info": info_habiles,
                "items_pendientes": faltan_compras_item,
                "valor_nacionalizacion": float(p.valor_nacionalizacion_dec),
                "total": float(p.total),
                "probabilidad": prob_p
            })

        m = metricas_por_asesor.get(p.asesor_id)
        if m:
            m["total_pos"] += 1
            m["volumen"] += float(p.total)
            m["nacionalizacion"] += float(p.valor_nacionalizacion_dec)
            m["prob_acum"] += prob_p
            m["dias_acum"] += dias_gest
            m["dias_count"] += 1
            m["items_pendientes_compra"] += faltan_compras_item
            if p.estado == "Entregado":
                m["entregados"] += 1
            else:
                m["pendientes"] += 1
            if p.plazo_entrega:
                m["sla_total"] += 1
                if en_tiempo:
                    m["sla_ok"] += 1

    probabilidad_global = round(suma_probabilidad / total_pedidos, 1) if total_pedidos else 0.0
    tiempo_promedio_dias = round(sum(tiempos_gestion) / len(tiempos_gestion), 1) if tiempos_gestion else 0.0
    cumplimiento_sla = round((entregas_a_tiempo / entregas_evaluadas) * 100, 1) if entregas_evaluadas else 100.0
    pendientes_gestion = sum(1 for p in pedidos if p.estado != "Entregado")

    tabla_asesores = []
    for aid, m in metricas_por_asesor.items():
        t_pos = m["total_pos"]
        m["tiempo_promedio"] = round(m["dias_acum"] / m["dias_count"], 1) if m["dias_count"] else 0.0
        m["probabilidad"] = round(m["prob_acum"] / t_pos, 1) if t_pos else 0.0
        m["sla_pct"] = round((m["sla_ok"] / m["sla_total"]) * 100, 1) if m["sla_total"] else 100.0
        score = (m["probabilidad"] * 0.6) + (m["sla_pct"] * 0.4) if t_pos else 0
        if t_pos == 0:
            m["nivel"] = "Sin casos"
        elif score >= 85:
            m["nivel"] = "Sobresaliente"
        elif score >= 65:
            m["nivel"] = "Óptimo"
        else:
            m["nivel"] = "En seguimiento"
        tabla_asesores.append(m)

    tabla_asesores.sort(key=lambda x: x["volumen"], reverse=True)
    todos_meses = sorted(list(set(ventas_mensuales_nac.keys()) | set(ventas_mensuales_int.keys())))
    top_prov_sorted = sorted(proveedores_top.items(), key=lambda x: x[1], reverse=True)[:6]

    return {
        "user": user,
        "pedidos_filtrados": pedidos,
        "todos_items_global": todos_items_global,
        "asesores_equipo": asesores_equipo,
        "f_asesor": asesor_filtro,
        "f_tipo": tipo_filtro,
        "f_desde": desde,
        "f_hasta": hasta,
        "total_pedidos": total_pedidos,
        "total_nacionales": total_nacionales,
        "total_internacionales": total_internacionales,
        "valor_total_ventas": valor_total_ventas,
        "valor_subtotal": valor_subtotal,
        "valor_total_nacionalizacion": valor_total_nacionalizacion,
        "pendientes_gestion": pendientes_gestion,
        "items_sin_compra": items_sin_compra,
        "probabilidad_global": probabilidad_global,
        "tiempo_promedio_dias": tiempo_promedio_dias,
        "cumplimiento_sla": cumplimiento_sla,
        "tabla_asesores": tabla_asesores,
        "compras_pendientes_po": compras_pendientes_po,
        "meses_labels": json.dumps(todos_meses),
        "meses_nac_data": json.dumps([round(ventas_mensuales_nac.get(m, 0), 2) for m in todos_meses]),
        "meses_int_data": json.dumps([round(ventas_mensuales_int.get(m, 0), 2) for m in todos_meses]),
        "estados_labels": json.dumps(list(estado_counts.keys())),
        "estados_data": json.dumps(list(estado_counts.values())),
        "asesores_labels": json.dumps([a["nombre"] for a in tabla_asesores]),
        "asesores_vol_data": json.dumps([round(a["volumen"], 2) for a in tabla_asesores]),
        "prov_labels": json.dumps([k for k, _ in top_prov_sorted]),
        "prov_data": json.dumps([round(v, 2) for _, v in top_prov_sorted]),
    }


@app.get("/informes", response_class=HTMLResponse)
def vista_informes(request: Request, asesor_id: str = "", tipo: str = "", desde: str = "", hasta: str = "",
                   user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    if user.rol not in ROLES_SUPERIORES:
        return RedirectResponse("/pedidos", status_code=303)
    ctx = _construir_analitica(db, user, asesor_id, tipo, desde, hasta)
    return render(request, "Informes.html", ctx)


@app.get("/informes/excel")
def descargar_informe_gerencial(asesor_id: str = "", tipo: str = "", desde: str = "", hasta: str = "",
                                user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    if user.rol not in ROLES_SUPERIORES:
        return RedirectResponse("/pedidos", status_code=303)

    ctx = _construir_analitica(db, user, asesor_id, tipo, desde, hasta)
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Metricas_Asesores"
    ws1.merge_cells("A1:K1")
    ws1["A1"] = "INFORME GERENCIAL DE DESEMPEÑO E INDICADORES OPERATIVOS - FLY LOGISTIC S.A.S."
    ws1["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws1["A1"].fill = PatternFill("solid", fgColor="002060")
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")

    enc1 = ["Asesor", "Cargo", "Superior", "Total POs", "Entregadas", "En Proceso",
            "Ítems Sin Compra", "Tiempo Prom. (Días)", "Cumplimiento (%)", "Nacionalización ($)", "Volumen Total ($)"]
    ws1.append(enc1)
    for c in ws1[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E40AF")

    for a in ctx["tabla_asesores"]:
        ws1.append([
            a["nombre"], a["rol"], a["jefe"], a["total_pos"], a["entregados"],
            a["pendientes"], a["items_pendientes_compra"], a["tiempo_promedio"],
            a["probabilidad"], a["nacionalizacion"], a["volumen"]
        ])
    for i in range(1, 12):
        ws1.column_dimensions[get_column_letter(i)].width = 20

    ws2 = wb.create_sheet("Seguimiento_POs")
    enc2 = ["PO", "Cliente", "Asesor", "Operación", "Estado", "Indicador Plazo",
            "Fecha Prometida", "Ítems Pendientes", "Nacionalización ($)", "Avance (%)", "Total PO ($)"]
    ws2.append(enc2)
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="002060")

    for cp in ctx["compras_pendientes_po"]:
        ws2.append([
            cp["po"], cp["cliente"], cp["asesor"], cp["tipo"], cp["estado"],
            cp["dias_habiles_info"]["etiqueta"], cp["plazo"], cp["items_pendientes"],
            cp["valor_nacionalizacion"], cp["probabilidad"], cp["total"]
        ])
    for i in range(1, 12):
        ws2.column_dimensions[get_column_letter(i)].width = 22

    ws3 = wb.create_sheet("Detalle_Global_Items")
    enc3 = ["PO", "Operación", "Asesor", "Cliente", "Proveedor", "Fecha Compra", "¿Comprado?",
            "Part Number", "Descripción", "Cantidad", "Unidad", "Estado Ítem", "Estado PO", "Indicador Plazo"]
    ws3.append(enc3)
    for c in ws3[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="002060")
    for it_g in ctx["todos_items_global"]:
        ws3.append([
            it_g["po"], it_g["tipo"], it_g["asesor"], it_g["cliente"], it_g["proveedor"],
            it_g["fecha_compra"], "SÍ" if it_g["comprado"] else "PENDIENTE",
            it_g["parte_numero"], it_g["descripcion"], it_g["cantidad"], it_g["unidad_medida"],
            it_g["estado_producto"], it_g["estado_po"], it_g["semaforo"]
        ])
    for i in range(1, 15):
        ws3.column_dimensions[get_column_letter(i)].width = 22

    carpeta_rep = os.path.join(BASE_DIR, "Reportes_Gerenciales")
    os.makedirs(carpeta_rep, exist_ok=True)
    ruta_rep = os.path.join(carpeta_rep, "Informe_Gerencial_FlyLogistic.xlsx")
    wb.save(ruta_rep)
    return FileResponse(ruta_rep, filename="Informe_Gerencial_FlyLogistic.xlsx", media_type=XLSX_MIME)


@app.get("/informes/productos_excel")
def descargar_excel_maestro_productos(asesor_id: str = "", tipo: str = "", estado_producto: str = "",
                                      user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    if user.rol not in ROLES_SUPERIORES:
        return RedirectResponse("/pedidos", status_code=303)

    visibles = ids_asesores_visibles(db, user)
    q_ped = db.query(models.Pedido).options(
        joinedload(models.Pedido.items),
        joinedload(models.Pedido.cliente),
        joinedload(models.Pedido.asesor),
        joinedload(models.Pedido.internacional)
    )
    if visibles is not None:
        q_ped = q_ped.filter(models.Pedido.asesor_id.in_(visibles))
    if asesor_id.isdigit():
        q_ped = q_ped.filter(models.Pedido.asesor_id == int(asesor_id))
    if tipo in ("Nacional", "Internacional"):
        q_ped = q_ped.filter(models.Pedido.tipo == tipo)
    if estado_producto in ESTADOS_PRODUCTO:
        q_ped = q_ped.filter(models.Pedido.items.any(models.ItemPedido.estado_producto == estado_producto))

    pedidos = q_ped.all()
    carpeta_rep = os.path.join(BASE_DIR, "Reportes_Gerenciales")
    os.makedirs(carpeta_rep, exist_ok=True)
    ruta_rep = os.path.join(carpeta_rep, "Control_Maestro_Productos_PO_FlyLogistic.xlsx")
    generar_excel_maestro_productos(ruta_rep, pedidos)
    return FileResponse(ruta_rep, filename="Control_Maestro_Productos_PO_FlyLogistic.xlsx", media_type=XLSX_MIME)


@app.get("/pedidos/consolidado/excel")
def descargar_consolidado_individual(asesor_id: str = "", user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    visibles = ids_asesores_visibles(db, user)
    objetivo_user = user

    if user.rol in ROLES_SUPERIORES and asesor_id.isdigit():
        cand = db.get(models.Asesor, int(asesor_id))
        if cand and (visibles is None or cand.id in visibles):
            objetivo_user = cand

    q_ped = db.query(models.Pedido).options(
        joinedload(models.Pedido.items),
        joinedload(models.Pedido.cliente),
        joinedload(models.Pedido.asesor),
        joinedload(models.Pedido.internacional)
    ).filter(models.Pedido.asesor_id == objetivo_user.id)

    pedidos_usuario = q_ped.all()
    carpeta_cons = os.path.join(BASE_DIR, "Consolidados_Usuarios")
    os.makedirs(carpeta_cons, exist_ok=True)
    nombre_arch = f"Consolidado_POs_{nombre_seguro(objetivo_user.nombre)}.xlsx"
    ruta_cons = os.path.join(carpeta_cons, nombre_arch)
    generar_excel_consolidado_usuario(ruta_cons, pedidos_usuario, objetivo_user.nombre)
    return FileResponse(ruta_cons, filename=nombre_arch, media_type=XLSX_MIME)


@app.get("/pedidos", response_class=HTMLResponse)
def mis_pedidos(request: Request, q: str = "", po: str = "", cliente: str = "", proveedor: str = "",
                asesor_id: str = "", estado: str = "", tipo: str = "", estado_producto: str = "",
                page: int = 1, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    visibles = ids_asesores_visibles(db, user)
    query = (db.query(models.Pedido)
             .options(
                 joinedload(models.Pedido.items),
                 joinedload(models.Pedido.cliente),
                 joinedload(models.Pedido.asesor),
                 joinedload(models.Pedido.internacional)
             )
             .join(models.Cliente))

    if visibles is not None:
        query = query.filter(models.Pedido.asesor_id.in_(visibles))
    if user.rol in ROLES_SUPERIORES and asesor_id.isdigit():
        query = query.filter(models.Pedido.asesor_id == int(asesor_id))
    if estado in ESTADOS:
        query = query.filter(models.Pedido.estado == estado)
    if tipo in ("Nacional", "Internacional"):
        query = query.filter(models.Pedido.tipo == tipo)
    if po.strip():
        query = query.filter(models.Pedido.po.ilike(f"%{po.strip()}%"))
    if cliente.strip():
        query = query.filter(models.Cliente.nombre.ilike(f"%{cliente.strip()}%"))
    if proveedor.strip():
        query = query.filter(models.Pedido.items.any(models.ItemPedido.proveedor.ilike(f"%{proveedor.strip()}%")))
    if estado_producto in ESTADOS_PRODUCTO:
        query = query.filter(models.Pedido.items.any(models.ItemPedido.estado_producto == estado_producto))
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(
            models.Pedido.po.ilike(like)
            | models.Cliente.nombre.ilike(like)
            | models.Pedido.items.any(models.ItemPedido.proveedor.ilike(like))
            | models.Pedido.items.any(models.ItemPedido.parte_numero.ilike(like))
        )

    pedidos_sin_ordenar = query.all()
    pedidos_ordenados = sorted(
        pedidos_sin_ordenar,
        key=lambda x: (
            1 if x.estado == "Entregado" else 0,
            x.dias_habiles_info["orden_urgencia"],
            -x.id
        )
    )

    por_pagina = 15
    total_pedidos = len(pedidos_ordenados)
    total_paginas = max(1, (total_pedidos + por_pagina - 1) // por_pagina)
    pagina_actual = max(1, min(page, total_paginas))

    inicio = (pagina_actual - 1) * por_pagina
    fin = inicio + por_pagina
    pedidos_paginados = pedidos_ordenados[inicio:fin]

    q_as = db.query(models.Asesor)
    if visibles is not None:
        q_as = q_as.filter(models.Asesor.id.in_(visibles))
    asesores_disponibles = q_as.order_by(models.Asesor.nombre).all()

    clientes_disponibles = [c[0] for c in db.query(models.Cliente.nombre).order_by(models.Cliente.nombre).all() if c[0]]
    proveedores_disponibles = sorted(list({
        r[0] for r in db.query(models.ItemPedido.proveedor).all() if r[0]
    }))

    return render(request, "pedidos.html", {
        "user": user,
        "pedidos": pedidos_paginados,
        "todos_pedidos": pedidos_ordenados,
        "pagina_actual": pagina_actual,
        "total_paginas": total_paginas,
        "total_pedidos": total_pedidos,
        "q": q,
        "po_f": po,
        "cliente_f": cliente,
        "proveedor_f": proveedor,
        "asesor_f": asesor_id,
        "estado_f": estado,
        "tipo_f": tipo,
        "estado_prod_f": estado_producto,
        "estados": ESTADOS,
        "estados_producto": ESTADOS_PRODUCTO,
        "asesores_disponibles": asesores_disponibles,
        "clientes_disponibles": clientes_disponibles,
        "proveedores_disponibles": proveedores_disponibles,
    })


@app.post("/pedidos/{pid}/notificar")
def notificar_po_responsable(pid: int, request: Request,
                             user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    if user.rol not in ROLES_SUPERIORES:
        request.session["flash"] = {
            "ok": False,
            "msg": "Acción reservada para Administración, Gerencia y Jefatura de Área."
        }
        return RedirectResponse("/pedidos", status_code=303)

    pedido = pedido_del_usuario(db, pid, user)
    if not pedido:
        request.session["flash"] = {"ok": False, "msg": "Orden no encontrada."}
        return RedirectResponse("/pedidos", status_code=303)
    destino = (pedido.asesor.email if (pedido.asesor and pedido.asesor.email) else None) or user.email or CORREO_OFICIAL

    asunto, cuerpo = construir_cuerpo_alerta_po(pedido)
    ok, detalle = enviar_correo_con_diagnostico(destino, asunto, cuerpo)
    if ok:
        request.session["flash"] = {
            "ok": True,
            "msg": f"Notificación enviada a {destino} ({pedido.dias_habiles_info['etiqueta']})."
        }
    else:
        request.session["flash"] = {
            "ok": False,
            "msg": f"Aviso SMTP: verifique las credenciales del servidor de correo en Mi Perfil ({detalle[:90]})."
        }
    return RedirectResponse("/pedidos", status_code=303)


@app.post("/pedidos/notificar_proximos")
def notificar_pos_proximas(request: Request, tareas: BackgroundTasks,
                           user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    if user.rol not in ROLES_SUPERIORES:
        request.session["flash"] = {
            "ok": False,
            "msg": "Acción reservada para Administración, Gerencia y Jefatura de Área."
        }
        return RedirectResponse("/pedidos", status_code=303)

    visibles = ids_asesores_visibles(db, user)
    q = db.query(models.Pedido).options(
        joinedload(models.Pedido.items),
        joinedload(models.Pedido.cliente),
        joinedload(models.Pedido.asesor),
        joinedload(models.Pedido.internacional)
    ).filter(models.Pedido.estado != "Entregado")
    if visibles is not None:
        q = q.filter(models.Pedido.asesor_id.in_(visibles))

    enviados = 0
    for p in q.all():
        destino = (p.asesor.email if (p.asesor and p.asesor.email) else None) or CORREO_OFICIAL
        asunto, cuerpo = construir_cuerpo_alerta_po(p)
        tareas.add_task(enviar_correo, destino, asunto, cuerpo)
        enviados += 1

    request.session["flash"] = {
        "ok": True,
        "msg": f"Se programó el envío de {enviados} notificaciones de seguimiento operativo."
    }
    return RedirectResponse("/pedidos", status_code=303)


@app.get("/api/validar_po")
def api_validar_po(po: str = "", excluir_id: int = 0, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    po_limpio = po.strip()
    if not po_limpio:
        return JSONResponse({"existe": False, "sugerido": sugerir_siguiente_po(db)})
    q = db.query(models.Pedido).filter(func.lower(models.Pedido.po) == po_limpio.lower())
    if excluir_id:
        q = q.filter(models.Pedido.id != excluir_id)
    existente = q.first()
    if existente:
        return JSONResponse({
            "existe": True,
            "po": existente.po,
            "tipo": existente.tipo,
            "asesor": existente.asesor.nombre if existente.asesor else "",
            "sugerido": sugerir_siguiente_po(db)
        })
    return JSONResponse({"existe": False, "sugerido": sugerir_siguiente_po(db)})


def _form(request, db: Session, user: models.Asesor, d, modo, pedido=None, error=None, status=200):
    visibles = ids_asesores_visibles(db, user)
    q_as = db.query(models.Asesor)
    if visibles is not None:
        q_as = q_as.filter(models.Asesor.id.in_(visibles))
    asesores_disponibles = q_as.order_by(models.Asesor.nombre).all()

    pos_existentes = [
        {"po": p.po, "tipo": p.tipo}
        for p in db.query(models.Pedido).order_by(models.Pedido.id.desc()).limit(30).all()
    ]

    return render(request, "pedido_form.html", {
        "user": user,
        "d": d,
        "modo": modo,
        "pedido": pedido,
        "error": error,
        "estados": ESTADOS,
        "estados_aduana": ESTADOS_ADUANA,
        "estados_producto": ESTADOS_PRODUCTO,
        "metodos_pago_opciones": METODOS_PAGO_OPCIONES,
        "unidades_medida": UNIDADES_MEDIDA,
        "asesores_disponibles": asesores_disponibles,
        "siguiente_po_sugerida": sugerir_siguiente_po(db),
        "pos_existentes": pos_existentes,
    }, status)


@app.get("/pedidos/nuevo", response_class=HTMLResponse)
def nuevo_form(request: Request, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    hoy = datetime.now().strftime("%Y-%m-%d")
    sugerida = sugerir_siguiente_po(db)
    d = {k: "" for k in CAMPOS}
    d.update(
        fecha_inicial=hoy,
        po_numero=sugerida,
        tipo_pedido="Nacional",
        iva_porcentaje="19",
        valor_nacionalizacion="0.00",
        llega_click_mail="No",
        estado_general="En proceso",
        asesor_asignado_id=str(user.id),
        items=[{
            "fecha_compra_prov": "", "proveedor": "", "parte_numero": "",
            "parte_numero_alterno": "", "unidad_medida": "UND", "descripcion": "",
            "cantidad": "1", "precio_unitario": "", "iva_item": "19",
            "estado_producto": "PENDIENTE"
        }]
    )
    return _form(request, db, user, d, "nuevo")


@app.post("/pedidos/nuevo", response_class=HTMLResponse)
async def nuevo_post(request: Request, tareas: BackgroundTasks,
                     user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    form_raw = await request.form()
    d = datos_desde_form(form_raw)
    archivo_doc = form_raw.get("archivo_documentacion")
    try:
        v = parsear(d)
        existente_po = db.query(models.Pedido).filter(func.lower(models.Pedido.po) == v["po"].lower()).first()
        if existente_po:
            raise ValueError(
                f"La orden «{v['po']}» ya se encuentra registrada ({existente_po.tipo}). "
                f"Consecutivo disponible: {sugerir_siguiente_po(db)}"
            )

        asesor_id_final = user.id
        if user.rol in ROLES_SUPERIORES and d.get("asesor_asignado_id", "").isdigit():
            asesor_id_final = int(d["asesor_asignado_id"])

        pedido = models.Pedido(po=v["po"], tipo=v["tipo"], asesor_id=asesor_id_final)
        aplicar(db, pedido, v)
        db.add(pedido)
        db.commit()
    except ValueError as e:
        db.rollback()
        return _form(request, db, user, d, "nuevo", error=str(e), status=400)
    except IntegrityError:
        db.rollback()
        return _form(request, db, user, d, "nuevo", error="El número de PO ya se encuentra registrado.", status=400)

    db.refresh(pedido)
    ruta_doc = None
    if isinstance(archivo_doc, UploadFile) and archivo_doc.filename:
        ruta_doc = await guardar_archivo_documentacion(pedido, archivo_doc)
        db.commit()
        db.refresh(pedido)

    estado = generar_salidas(db, pedido, ruta_doc)

    if pedido.plazo_entrega:
        destino = (pedido.asesor.email if (pedido.asesor and pedido.asesor.email) else None) or CORREO_OFICIAL
        asunto, cuerpo = construir_cuerpo_alerta_po(pedido)
        tareas.add_task(enviar_correo, destino, asunto, cuerpo)

    request.session["flash"] = {"ok": True, "msg": f"Orden {pedido.po_codigo} registrada. {estado}."}
    return RedirectResponse("/pedidos", status_code=303)


@app.get("/pedidos/{pid}/editar", response_class=HTMLResponse)
def editar_form(pid: int, request: Request, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    pedido = pedido_del_usuario(db, pid, user)
    if not pedido:
        request.session["flash"] = {"ok": False, "msg": "Orden no encontrada o sin permisos de acceso."}
        return RedirectResponse("/pedidos", status_code=303)
    return _form(request, db, user, datos_desde_pedido(pedido), "editar", pedido)


@app.post("/pedidos/{pid}/editar", response_class=HTMLResponse)
async def editar_post(pid: int, request: Request,
                      user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    pedido = pedido_del_usuario(db, pid, user)
    if not pedido:
        request.session["flash"] = {"ok": False, "msg": "Orden no encontrada."}
        return RedirectResponse("/pedidos", status_code=303)
    form_raw = await request.form()
    d = datos_desde_form(form_raw)
    archivo_doc = form_raw.get("archivo_documentacion")
    d["po_numero"], d["tipo_pedido"] = pedido.po, pedido.tipo
    try:
        aplicar(db, pedido, parsear(d, tipo_fijo=pedido.tipo))
        if user.rol in ROLES_SUPERIORES and d.get("asesor_asignado_id", "").isdigit():
            pedido.asesor_id = int(d["asesor_asignado_id"])
        db.commit()
    except ValueError as e:
        db.rollback()
        return _form(request, db, user, d, "editar", pedido, error=str(e), status=400)
    except IntegrityError:
        db.rollback()
        return _form(request, db, user, d, "editar", pedido, error="Error de integridad al actualizar la orden.", status=400)

    db.refresh(pedido)
    ruta_doc = None
    if isinstance(archivo_doc, UploadFile) and archivo_doc.filename:
        ruta_doc = await guardar_archivo_documentacion(pedido, archivo_doc)
        db.commit()
        db.refresh(pedido)

    estado = generar_salidas(db, pedido, ruta_doc)
    request.session["flash"] = {"ok": True, "msg": f"Orden {pedido.po_codigo} actualizada. {estado}."}
    return RedirectResponse("/pedidos", status_code=303)


@app.get("/pedidos/{pid}/documento")
def descargar_documento_adjunto(pid: int, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    pedido = pedido_del_usuario(db, pid, user)
    if not pedido or not pedido.documentacion_archivo:
        return RedirectResponse("/pedidos", status_code=303)
    ruta_doc = os.path.join(carpeta_po_local(pedido), pedido.documentacion_archivo)
    if not os.path.exists(ruta_doc):
        return RedirectResponse("/pedidos", status_code=303)
    return FileResponse(ruta_doc, filename=pedido.documentacion_archivo)


@app.post("/pedidos/{pid}/eliminar")
def eliminar_pedido(pid: int, request: Request, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    if user.rol not in ROLES_SUPERIORES:
        request.session["flash"] = {"ok": False, "msg": "Permisos insuficientes para eliminar órdenes."}
        return RedirectResponse("/pedidos", status_code=303)
    pedido = pedido_del_usuario(db, pid, user)
    if not pedido:
        request.session["flash"] = {"ok": False, "msg": "Orden no encontrada."}
        return RedirectResponse("/pedidos", status_code=303)
    po_num = pedido.po_codigo
    db.delete(pedido)
    db.commit()
    request.session["flash"] = {"ok": True, "msg": f"Orden {po_num} eliminada del sistema."}
    return RedirectResponse("/pedidos", status_code=303)


@app.get("/pedidos/{pid}/excel")
def descargar_excel(pid: int, user: models.Asesor = Depends(usuario_actual), db: Session = Depends(get_db)):
    pedido = pedido_del_usuario(db, pid, user)
    if not pedido:
        return RedirectResponse("/pedidos", status_code=303)
    ruta, nombre = ruta_excel(pedido)
    generar_excel_com_fm_08(pedido, ruta)
    return FileResponse(ruta, filename=nombre, media_type=XLSX_MIME)


@app.api_route("/logo-empresa", methods=["GET", "HEAD"])
def servir_logo_corporativo():
    candidatos = [
        os.path.join(BASE_DIR, "logo_fly.webp.webp"),
        os.path.join(BASE_DIR, "logo_fly.webp"),
        os.path.join(BASE_DIR, "cropped-Logo-Fly-Logistic-1-1-convertido-de-png-600x233.webp"),
        os.path.join(BASE_DIR, "logo_fly.png"),
        os.path.join(BASE_DIR, "logo.png"),
        "logo_fly.webp.webp",
        "logo_fly.webp",
    ]
    for ruta in candidatos:
        if os.path.exists(ruta):
            mime = "image/webp" if ruta.lower().endswith(".webp") else "image/png"
            return FileResponse(ruta, media_type=mime)
    svg_fallback = (
        '<svg xmlns="http://www.w3.org/2000/svg" width="180" height="48" viewBox="0 0 180 48">'
        '<rect width="180" height="48" rx="8" fill="#002060"/>'
        '<text x="90" y="30" font-family="Arial,sans-serif" font-size="18" font-weight="bold" fill="#ffffff" text-anchor="middle">FLY LOGISTIC</text>'
        '</svg>'
    )
    return HTMLResponse(content=svg_fallback, media_type="image/svg+xml")
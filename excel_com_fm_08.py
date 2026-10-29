import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

COLORES_FILAS_PO = ["FFFF00", "92D050", "B8CCE4", "FCD5B4", "E5E0EC", "D8E4BC"]

ENC_NACIONAL = [
    "1. Fecha inicial",
    "2. PO (Orden de compra)",
    "3. Cliente",
    "4. Nombre del asesor",
    "5. Fecha de compra a proveedores",
    "6. Proveedor (Origen)",
    "7. Parte número / Part Number",
    "8. Parte número alterno",
    "9. Descripción del producto",
    "10. Unidad de medida",
    "11. Cantidad",
    "12. Estado del Producto",
    "13. Método de pago",
    "14. Precio unitario inicial ($)",
    "15. Precio total ($)",
    "16. IVA (19%) (%)",
    "17. Valor después de IVA ($)",
    "18. Plazo de entrega prometido",
    "19. Días Hábiles (Semáforo)",
    "20. Fecha de entrega al cliente",
    "21. Número de remisión",
    "22. Documentación",
    "23. Estado PO",
    "24. Observaciones",
]

ENC_INTERNACIONAL = [
    "1. Fecha inicial",
    "2. PO (Orden de compra)",
    "3. Cliente",
    "4. Nombre del asesor",
    "5. Fecha de compra a proveedores",
    "6. Proveedor (Origen)",
    "7. Parte número / Part Number",
    "8. Parte número alterno",
    "9. Descripción del producto",
    "10. Unidad de medida",
    "11. Cantidad",
    "12. Estado del Producto",
    "13. Método de pago",
    "14. Precio unitario inicial ($)",
    "15. Precio total ($)",
    "16. TAX (7%) (%)",
    "17. Valor después de TAX + Nac. ($)",
    "18. Plazo de entrega prometido",
    "19. Días Hábiles (Semáforo)",
    "20. Fecha de entrega warehouse EEUU",
    "21. Tracking",
    "22. Nacionalización (Estado)",
    "23. Valor de la nacionalización ($)",
    "24. Fecha de Nacionalización",
    "25. Documentación",
    "26. Estado PO",
    "27. Observaciones",
]


def _escribir_bloque_po_en_hoja(ws, p, fila_inicio: int, color_hex: str) -> int:
    es_nacional = (p.tipo == "Nacional")
    intl = p.internacional
    val_nac_total = float(p.valor_nacionalizacion or 0)
    items_lista = list(p.items)
    num_items = max(len(items_lista), 1)
    val_nac_por_item = round(val_nac_total / num_items, 2) if not es_nacional else 0.0
    semaforo = p.dias_habiles_info["detalle"]
    doc_texto = p.documentacion or ""
    if getattr(p, "documentacion_archivo", None) and p.documentacion_archivo not in doc_texto:
        doc_texto = f"{doc_texto} [{p.documentacion_archivo}]".strip()

    relleno_fila = PatternFill("solid", fgColor=color_hex)
    borde_fino = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    fila_actual = fila_inicio
    for item in items_lista:
        iva_pct = float(item.iva_porcentaje if item.iva_porcentaje is not None else (19 if es_nacional else 7))
        base_fila = [
            p.fecha_inicial,
            p.po,
            p.cliente.nombre if p.cliente else "",
            p.asesor.nombre if p.asesor else "",
            item.fecha_compra_proveedor,
            item.proveedor,
            item.parte_numero,
            item.parte_numero_alterno or "",
            item.descripcion,
            item.unidad_medida or "UND",
            int(item.cantidad),
            getattr(item, "estado_producto", "PENDIENTE") or "PENDIENTE",
            p.metodo_pago or "",
            float(item.precio_unitario),
            f"=K{fila_actual}*N{fila_actual}",
            iva_pct,
        ]

        if es_nacional:
            valor_despues = f"=O{fila_actual}*(1+P{fila_actual}/100)"
            cola = [
                valor_despues,
                p.plazo_entrega,
                semaforo,
                p.fecha_entrega_cliente,
                p.numero_remision or "",
                doc_texto,
                p.estado,
                p.observaciones or "",
            ]
            ws.append(base_fila + cola)
        else:
            valor_despues = f"=O{fila_actual}*(1+P{fila_actual}/100)+W{fila_actual}"
            cola = [
                valor_despues,
                p.plazo_entrega,
                semaforo,
                intl.fecha_click_mail if intl else None,
                (intl.tracking or "") if intl else "",
                (intl.estado_nacionalizacion or "No aplica / Pendiente") if intl else "No aplica / Pendiente",
                val_nac_por_item,
                intl.fecha_nacionalizacion if intl else None,
                doc_texto,
                p.estado,
                p.observaciones or "",
            ]
            ws.append(base_fila + cola)

        total_cols = 24 if es_nacional else 27
        for col_i in range(1, total_cols + 1):
            celda = ws.cell(fila_actual, col_i)
            celda.fill = relleno_fila
            celda.border = borde_fino
            if hasattr(celda.value, "year"):
                celda.number_format = "YYYY-MM-DD"
        for col_moneda in ("N", "O", "Q"):
            ws[f"{col_moneda}{fila_actual}"].number_format = '"$"#,##0.00'
        if not es_nacional:
            ws[f"W{fila_actual}"].number_format = '"$"#,##0.00'

        fila_actual += 1

    ult_item_fila = fila_actual - 1
    ws[f"M{fila_actual}"] = f"TOTALES PO {p.po}:"
    ws[f"M{fila_actual}"].font = Font(bold=True, color="002060")
    ws[f"O{fila_actual}"] = f"=SUM(O{fila_inicio}:O{ult_item_fila})"
    ws[f"O{fila_actual}"].font = Font(bold=True)
    ws[f"O{fila_actual}"].number_format = '"$"#,##0.00'
    ws[f"Q{fila_actual}"] = f"=SUM(Q{fila_inicio}:Q{ult_item_fila})"
    ws[f"Q{fila_actual}"].font = Font(bold=True, color="002060")
    ws[f"Q{fila_actual}"].number_format = '"$"#,##0.00'

    if not es_nacional:
        ws[f"W{fila_actual}"] = f"=SUM(W{fila_inicio}:W{ult_item_fila})"
        ws[f"W{fila_actual}"].font = Font(bold=True, color="002060")
        ws[f"W{fila_actual}"].number_format = '"$"#,##0.00'

    total_cols = 24 if es_nacional else 27
    for col_i in range(1, total_cols + 1):
        c_tot = ws.cell(fila_actual, col_i)
        c_tot.fill = relleno_fila
        c_tot.border = borde_fino

    return fila_actual + 1


def _preparar_hoja_cuadro(ws, es_nacional: bool):
    if es_nacional:
        ws.merge_cells("A1:X1")
        titulo, encabezados = "CUADRO DE COMPRAS NACIONALES FLY LOGISTIC 2026", ENC_NACIONAL
    else:
        ws.merge_cells("A1:AA1")
        titulo, encabezados = "CUADRO DE COMPRAS EN EL EXTERIOR FLY LOGISTIC 2026", ENC_INTERNACIONAL

    t = ws["A1"]
    t.value = titulo
    t.font = Font(bold=True, size=14, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor="002060")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.append(encabezados)
    relleno_enc = PatternFill("solid", fgColor="2563EB")
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = relleno_enc
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 30

    for i, enc in enumerate(encabezados, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(18, len(enc) * 0.95)


def generar_excel_com_fm_08(pedido, ruta_salida):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FLY2026"
    es_nacional = (pedido.tipo == "Nacional")
    _preparar_hoja_cuadro(ws, es_nacional)
    _escribir_bloque_po_en_hoja(ws, pedido, 3, "FFFF00")
    os.makedirs(os.path.dirname(os.path.abspath(ruta_salida)), exist_ok=True)
    wb.save(ruta_salida)
    return ruta_salida


def generar_excel_consolidado_usuario(ruta: str, pedidos_usuario: list, nombre_usuario: str):
    pedidos_ordenados = sorted(
        pedidos_usuario,
        key=lambda x: (
            1 if x.estado in ("Entregado", "Cancelado") else 0,
            x.dias_habiles_info["orden_urgencia"],
            x.plazo_entrega or datetime(2099, 12, 31).date()
        )
    )
    nacionales = [p for p in pedidos_ordenados if p.tipo == "Nacional"]
    internacionales = [p for p in pedidos_ordenados if p.tipo == "Internacional"]

    wb = openpyxl.Workbook()
    ws_nac = wb.active
    ws_nac.title = "Consolidado_Nacional"
    _preparar_hoja_cuadro(ws_nac, True)
    fila_n = 3
    for idx_po, p in enumerate(nacionales):
        color = COLORES_FILAS_PO[idx_po % len(COLORES_FILAS_PO)]
        fila_n = _escribir_bloque_po_en_hoja(ws_nac, p, fila_n, color)

    ws_int = wb.create_sheet("Consolidado_Internacional")
    _preparar_hoja_cuadro(ws_int, False)
    fila_i = 3
    for idx_po, p in enumerate(internacionales):
        color = COLORES_FILAS_PO[idx_po % len(COLORES_FILAS_PO)]
        fila_i = _escribir_bloque_po_en_hoja(ws_int, p, fila_i, color)

    os.makedirs(os.path.dirname(os.path.abspath(ruta)), exist_ok=True)
    wb.save(ruta)


def generar_excel_maestro_productos(ruta: str, pedidos: list):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estado_Global_Productos"

    ws.merge_cells("A1:P1")
    ws["A1"] = "CONTROL MAESTRO DE PRODUCTOS, COMPRAS Y ESTADOS POR PO · FLY LOGISTIC 2026"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="002060")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    enc = [
        "PO", "Tipo Operación", "Asesor Responsable", "Cliente", "Fecha Inicial PO",
        "Plazo Entrega PO", "Semáforo PO", "Estado General PO", "Proveedor (Origen)",
        "Fecha Compra Proveedor", "¿Comprado?", "Part Number / Parte Número", "Descripción Producto",
        "Cantidad", "Unidad", "Estado del Producto"
    ]
    ws.append(enc)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="1E40AF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 28

    pedidos_urgencia = sorted(
        pedidos,
        key=lambda x: (
            1 if x.estado in ("Entregado", "Cancelado") else 0,
            x.dias_habiles_info["orden_urgencia"]
        )
    )

    fila_idx = 3
    for p in pedidos_urgencia:
        for it in p.items:
            comprado_txt = "SÍ" if it.fecha_compra_proveedor else "FALTA COMPRAR"
            sku = f"{it.parte_numero} (Alt: {it.parte_numero_alterno})" if it.parte_numero_alterno else it.parte_numero
            ws.append([
                p.po,
                p.tipo,
                p.asesor.nombre if p.asesor else "",
                p.cliente.nombre if p.cliente else "",
                p.fecha_inicial.isoformat() if p.fecha_inicial else "",
                p.plazo_entrega.isoformat() if p.plazo_entrega else "Sin definir",
                p.dias_habiles_info["etiqueta"],
                p.estado,
                it.proveedor,
                it.fecha_compra_proveedor.isoformat() if it.fecha_compra_proveedor else "PENDIENTE COMPRA",
                comprado_txt,
                sku,
                it.descripcion,
                int(it.cantidad),
                it.unidad_medida or "UND",
                it.estado_producto or "PENDIENTE",
            ])
            color_est = "FEE2E2" if not it.fecha_compra_proveedor or (it.estado_producto == "PENDIENTE") else "DCFCE7"
            ws[f"K{fila_idx}"].fill = PatternFill("solid", fgColor=color_est)
            ws[f"P{fila_idx}"].fill = PatternFill("solid", fgColor=color_est)
            fila_idx += 1

    for i in range(1, len(enc) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22

    os.makedirs(os.path.dirname(os.path.abspath(ruta)), exist_ok=True)
    wb.save(ruta)